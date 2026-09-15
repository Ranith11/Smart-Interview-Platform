import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    return heading

def generate_report():
    doc = Document()

    # Title
    title = doc.add_heading('SmartInterview: End-to-End System Workflow', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("A comprehensive guide to the architecture and data processing pipeline of the SmartInterview project.")
    
    # Section 1: Introduction
    add_heading(doc, '1. Project Overview')
    doc.add_paragraph(
        "SmartInterview is an AI-powered technical mock interview platform. "
        "It takes a candidate's resume, parses it for skills and project experience, "
        "and generates hyper-relevant, domain-specific interview questions. To ensure "
        "the questions are technically accurate and context-rich, it utilizes a "
        "Retrieval-Augmented Generation (RAG) pipeline backed by a custom technical knowledge base."
    )

    # Section 2: Data Collection
    add_heading(doc, '2. Data Collection (Web Scraping)')
    doc.add_paragraph(
        "The first step of the pipeline involves building the raw knowledge base."
    )
    p = doc.add_paragraph()
    p.add_run("Process: ").bold = True
    p.add_run("The system scrapes high-quality technical articles, tutorials, and interview prep content from authoritative sources (like GeeksforGeeks). Scripts like download_source.py and download_dsa_sources.py handle this task.")
    p2 = doc.add_paragraph()
    p2.add_run("Technologies Used: ").bold = True
    p2.add_run("Python, BeautifulSoup (bs4), Requests.")
    
    # Section 3: Data Cleaning & Processing
    add_heading(doc, '3. Data Cleaning and Chunking')
    doc.add_paragraph(
        "Raw HTML data contains boilerplate, navigation bars, ads, and irrelevant content that would confuse the AI."
    )
    p = doc.add_paragraph()
    p.add_run("Process: ").bold = True
    p.add_run("Using extract_text.py and clean_text.py, the raw HTML is stripped down to its core semantic content. Once the text is clean, chunk_text.py splits large articles into smaller, token-optimized 'chunks'. Each chunk is tagged with its domain (e.g., dbms, dsa, oop) and specific concept (e.g., indexing, binary trees) and saved as JSON files in the data/processed/ directory.")
    
    # Section 4: Vector Database (RAG Backend)
    add_heading(doc, '4. Knowledge Base Creation (Vector DB)')
    doc.add_paragraph(
        "To enable the AI to retrieve this knowledge in real-time during question generation, the text chunks must be converted into numerical vectors (embeddings)."
    )
    p = doc.add_paragraph()
    p.add_run("Process: ").bold = True
    p.add_run("The script ingest_vector_db.py reads the processed JSON chunks. It uses an embedding model to convert the text into dense vectors, and upserts them into a local ChromaDB instance (stored in chroma_db/). It also attaches the domain and concept metadata so queries can be filtered by domain.")
    p2 = doc.add_paragraph()
    p2.add_run("Technologies Used: ").bold = True
    p2.add_run("SentenceTransformers (Model: all-MiniLM-L6-v2), ChromaDB.")

    # Section 5: Resume Parsing
    add_heading(doc, '5. Candidate Resume Parsing')
    doc.add_paragraph(
        "When an interview starts, the system analyzes the candidate's background."
    )
    p = doc.add_paragraph()
    p.add_run("Process: ").bold = True
    p.add_run("The parse_resume.py script reads the candidate's PDF resume. It extracts the raw text and parses out their specific technical skills (e.g., Java, React, SQL) and their project experience. This forms the blueprint for the interview.")
    p2 = doc.add_paragraph()
    p2.add_run("Technologies Used: ").bold = True
    p2.add_run("PyMuPDF (fitz).")

    # Section 6: Question Generation (LLM + RAG)
    add_heading(doc, '6. Question Generation Pipeline')
    doc.add_paragraph(
        "The core intelligence of the system lies in generate_question.py. This script coordinates the LLM, the vector database, and the resume data to generate mock interview questions."
    )
    
    doc.add_paragraph("The workflow executes in several steps:", style='List Bullet')
    
    p = doc.add_paragraph(style='List Bullet 2')
    p.add_run("Skill Selection: ").bold = True
    p.add_run("The system picks a skill from the resume (e.g., 'SQL (Postgres)').")
    
    p = doc.add_paragraph(style='List Bullet 2')
    p.add_run("Domain Resolution: ").bold = True
    p.add_run("It maps the skill to its technical domain (e.g., SQL -> dbms) using a hardcoded Skill-Domain map.")
    
    p = doc.add_paragraph(style='List Bullet 2')
    p.add_run("RAG Retrieval: ").bold = True
    p.add_run("It constructs a highly specific semantic query (e.g., 'SQL database concepts: indexing, joins, normalization') and queries ChromaDB. ChromaDB returns the top 3 most relevant knowledge chunks for that exact domain.")
    
    p = doc.add_paragraph(style='List Bullet 2')
    p.add_run("Prompt Construction: ").bold = True
    p.add_run("It builds a prompt containing the candidate's project context, the retrieved RAG knowledge, the target skill, and instructions for difficulty/type (e.g., 'Hard', 'Scenario-based').")
    
    p = doc.add_paragraph(style='List Bullet 2')
    p.add_run("LLM Generation: ").bold = True
    p.add_run("The prompt is sent to the LLM via the Groq API. The script includes robust error handling, exponential backoff for API limits, and truncation detection to ensure questions aren't cut off mid-sentence (expanding max_tokens dynamically up to 1024 tokens if needed).")

    p2 = doc.add_paragraph()
    p2.add_run("Technologies Used: ").bold = True
    p2.add_run("Groq API (Model: openai/gpt-oss-120b), python-dotenv.")

    # Section 7: Summary of the Tech Stack
    add_heading(doc, '7. Full Technology Stack Overview')
    doc.add_paragraph("Language: Python", style='List Bullet')
    doc.add_paragraph("Data Scraping: BeautifulSoup4, requests", style='List Bullet')
    doc.add_paragraph("PDF Processing: PyMuPDF", style='List Bullet')
    doc.add_paragraph("Embeddings: SentenceTransformers (all-MiniLM-L6-v2)", style='List Bullet')
    doc.add_paragraph("Vector Database: ChromaDB", style='List Bullet')
    doc.add_paragraph("LLM Provider: Groq API", style='List Bullet')

    # Save Document
    file_path = os.path.join(os.getcwd(), 'SmartInterview_Workflow.docx')
    doc.save(file_path)
    print(f"Document successfully created at: {file_path}")

if __name__ == "__main__":
    generate_report()
