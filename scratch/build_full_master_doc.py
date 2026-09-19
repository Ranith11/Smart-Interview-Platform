import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_master_document():
    doc = Document()

    # Page setup - 0.75 inch margins for professional report
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill_hex)
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_title(text, subtitle=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

        if subtitle:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(2)
            p2.paragraph_format.space_after = Pt(14)
            run2 = p2.add_run(subtitle)
            run2.font.name = 'Arial'
            run2.font.size = Pt(11)
            run2.font.italic = True
            run2.font.color.rgb = RGBColor(79, 70, 229)

    def add_heading_styled(text, level):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(14 if level == 1 else (10 if level == 2 else 6))
        h.paragraph_format.space_after = Pt(3)
        run = h.runs[0]
        run.font.name = 'Arial'
        if level == 1:
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif level == 2:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(79, 70, 229)
        elif level == 3:
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
        return h

    def add_p(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Arial'
            r_pre.font.size = Pt(9.5)
            r_pre.font.color.rgb = RGBColor(30, 41, 59)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(9.5)
        run.italic = italic
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_bullet(text, bold_prefix=None, level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Arial'
            r_pre.font.size = Pt(9.5)
            r_pre.font.color.rgb = RGBColor(30, 41, 59)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_code_block(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.0)
        run.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_table_data(headers, data, col_widths=None):
        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        hdr_cells = table.rows[0].cells
        for i, title in enumerate(headers):
            hdr_cells[i].text = title
            set_cell_background(hdr_cells[i], '1E293B')
            set_cell_margins(hdr_cells[i], top=90, bottom=90, left=120, right=120)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = 'Arial'

        for row_idx, row_data in enumerate(data):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = 'F8FAFC' if row_idx % 2 == 1 else 'FFFFFF'
            for col_idx, cell_value in enumerate(row_data):
                row_cells[col_idx].text = str(cell_value)
                set_cell_background(row_cells[col_idx], bg_color)
                set_cell_margins(row_cells[col_idx], top=70, bottom=70, left=120, right=120)
                p = row_cells[col_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(8.0)
                    run.font.color.rgb = RGBColor(51, 65, 85)
                    run.font.name = 'Arial'

        if col_widths:
            for row in table.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Inches(w)

        doc.add_paragraph()

    # DOCUMENT TITLE
    add_title("SmartInterview: Master Technical Documentation", "Complete End-to-End System Architecture, Codebase Audit & Engineering Reference")

    # ============================================================
    # SECTION 1
    # ============================================================
    add_heading_styled("1. Project Overview", 1)
    add_p("SmartInterview is an intelligent, full-stack mock technical interview platform built to simulate real-world software engineering interviews. Traditional preparation platforms rely on static question banks or expensive human mock interviewers. SmartInterview bridges this gap by combining automated resume parsing, targeted job description mapping, syllabus-grounded academic assessment, domain-filtered Retrieval-Augmented Generation (RAG), dynamic question generation via high-speed LLMs, deterministic adaptive question sequencing based on Bloom's Taxonomy, and multi-dimensional automated answer evaluation.")
    
    add_heading_styled("Core Problems Solved", 2)
    add_bullet(" Solves the problem of generic interview prep by extracting verified skills and real personal project descriptions directly from PDF resumes.", "Personalized Questioning:")
    add_bullet(" Solves preparation misalignment by cross-matching candidate resume skills against target Job Description requirements into high-priority overlap skills (Group A) and critical role gaps (Group B).", "Job Description Alignment:")
    add_bullet(" Solves the limitation of static question lists by using a deterministic Python engine that dynamically modulates question difficulty and cognitive depth (Bloom's Taxonomy) after every single response.", "Dynamic Adaptive Assessment:")
    add_bullet(" Eliminates hallucinated questions by retrieving vetted computer science knowledge chunks from an isolated vector database (ChromaDB) before generating each question.", "RAG Grounding:")
    add_bullet(" Replaces arbitrary pass/fail grading with a 5-dimension rubric combining LLM technical critique, SBERT cosine semantic similarity, and algorithmic concept coverage.", "Holistic Multi-Dimensional Evaluation:")
    add_bullet(" Enables university students and course participants to upload course materials (PDF, TXT, DOCX) to practice against specific course syllabi in an isolated temporary vector space.", "Academic Syllabus Mode:")

    add_heading_styled("Comprehensive Technology Stack", 2)
    tech_stack = [
        ["Layer", "Technologies", "Actual Implementation Detail"],
        ["Frontend UI", "React 18, Vite 8, Tailwind CSS v4", "Single-Page Application in frontend/src/, configured on port 5174 with reverse-proxy to :8000."],
        ["Routing & State", "React Router v7, Context API", "App.jsx routing, AuthContext.jsx managing JWT in localStorage and /api/auth/me token refresh."],
        ["HTTP Client", "Axios 1.20.0", "api.js interceptors injecting Bearer token on requests and redirecting on 401 Unauthorized responses."],
        ["Icons & Design", "Lucide React 1.38", "Tailwind CSS modern glassmorphic dashboard, responsive sidebar, metric bars, and progress rings."],
        ["Backend API", "FastAPI 0.115+, Uvicorn", "Asynchronous Python web framework in backend/app/, auto OpenAPI docs at /docs, CORS middleware."],
        ["Database", "MySQL 8.0, PyMySQL, SQLAlchemy 2.0", "Database smartinterview with 7 core relational tables, foreign keys, cascade deletes, JSON columns."],
        ["Authentication", "JWT (python-jose HS256), bcrypt", "Stateless bearer token authentication, bcrypt password hashing truncated to 72-byte max."],
        ["LLM Provider", "Groq API (openai/gpt-oss-120b)", "High-throughput cloud inference engine with exponential retry backoff and token budget control."],
        ["RAG Vector DB", "ChromaDB (PersistentClient)", "Local vector database located in chroma_db/, collection technical_kb with cosine distance."],
        ["Embeddings", "SentenceTransformer all-MiniLM-L6-v2", "Dense vector embeddings (384 dimensions), local inference without external API cost."],
        ["Document Parsing", "PyMuPDF (fitz), python-docx", "Robust PDF text extraction with generic page header, footer, and pagination filtering."],
        ["Adaptive Engine", "Deterministic Python Algorithms", "Bloom's Taxonomy (Levels 1-6) + difficulty (easy/medium/hard) + weakness-biased round robin."],
        ["Evaluation", "Hybrid LLM + SBERT + Concept Math", "30% technical, 20% completeness, 20% relevance, 15% semantic similarity, 15% concept coverage."]
    ]
    add_table_data(["System Layer", "Technology", "Implementation Reference"], tech_stack[1:], [1.4, 2.3, 3.3])

    # ============================================================
    # SECTION 2
    # ============================================================
    add_heading_styled("2. Complete Project Architecture", 1)
    add_p("The SmartInterview architecture connects the browser client to an asynchronous FastAPI service layer, which coordinates document ingestion, vector retrieval, LLM generation, deterministic adaptive logic, and relational persistence.")

    arch_diagram = (
        "+---------------------------------------------------------------------------------------------------+\n"
        "|                                     REACT FRONTEND (Vite / Port 5174)                             |\n"
        "|  Landing -> Login/Register -> Dashboard -> Resume/JD Upload -> Setup -> Interview -> Results/Hist  |\n"
        "+---------------------------------------------------------------------------------------------------+\n"
        "                                                |  HTTP / REST (Axios with Bearer JWT)\n"
        "                                                v\n"
        "+---------------------------------------------------------------------------------------------------+\n"
        "|                                    FASTAPI BACKEND (Uvicorn / Port 8000)                          |\n"
        "|  Routers: /api/auth | /api/resumes | /api/job-descriptions | /api/interviews | /api/users          |\n"
        "+---------------------------------------------------------------------------------------------------+\n"
        "      |                          |                             |                        |\n"
        "      | Password / Token         | PDF Extraction              | Question Loop          | Persistence\n"
        "      v                          v                             v                        v\n"
        "+--------------+      +---------------------+      +----------------------+   +---------------------+\n"
        "| Auth Service |      | Resume & JD Service |      |  Interview Service   |   |   MySQL 8 Database  |\n"
        "| - bcrypt hash|      | - PyMuPDF (fitz)    |      |  - submit_and_eval() |   | - users             |\n"
        "| - JWT HS256  |      | - Section Regex     |      |  - Bloom Progression |   | - resumes           |\n"
        "+--------------+      | - Skill Taxonomy    |      |  - Skill Selection   |   | - job_descriptions  |\n"
        "                      +---------------------+      +----------------------+   | - interview_sessions|\n"
        "                                 |                             |              | - interview_questions\n"
        "                                 v (Matched Skills)            |              | - answers           |\n"
        "                      +----------------------------------------+              | - answer_evaluations|\n"
        "                      |                                                       +---------------------+\n"
        "                      v\n"
        "+---------------------------------------------------------------------------------------------------+\n"
        "|                                     QUESTION GENERATION PIPELINE                                  |\n"
        "| 1. SkillDomainMap: Canonical Skill -> Relevant KB Domain (e.g. 'sql' -> 'dbms')                    |\n"
        "| 2. DomainQueryTemplate: Construct domain-aware query string                                       |\n"
        "| 3. ChromaDB Query: SentenceTransformer all-MiniLM-L6-v2 -> retrieve top-3 chunks filtered by domain |\n"
        "| 4. Prompt Assembly: System prompt + Bloom instruction + RAG chunks + Project context + Dedup     |\n"
        "| 5. Groq LLM Inference: openai/gpt-oss-120b call with retry/backoff -> Single Clean Question Text  |\n"
        "+---------------------------------------------------------------------------------------------------+\n"
        "                                                |\n"
        "                                                v (Candidate Submits Text Answer)\n"
        "+---------------------------------------------------------------------------------------------------+\n"
        "|                                      ANSWER EVALUATION PIPELINE                                   |\n"
        "| 1. Groq LLM Critique: Technical (30%), Completeness (20%), Relevance (20%), Feedback, Strengths   |\n"
        "| 2. SentenceTransformer: Cosine similarity between candidate answer and (Question + Reference RAG) |\n"
        "| 3. Concept Coverage Math: Average of LLM score and len(found_concepts) / len(expected_concepts)  |\n"
        "| 4. Weighted Overall Score Calculation (Clamped 0-100) -> Persist AnswerEvaluation in MySQL        |\n"
        "| 5. Adaptive Decision: Score updates SkillPerformance -> Decide next Bloom level & next skill      |\n"
        "+---------------------------------------------------------------------------------------------------+"
    )
    add_code_block(arch_diagram)

    add_heading_styled("Detailed Subsystem Interconnections", 2)
    add_bullet(" Sends credentials to POST /api/auth/login. auth_service.py verifies bcrypt hash against MySQL users table and generates a signed JWT token returned to frontend localStorage.", "Frontend to Auth Router:")
    add_bullet(" User uploads PDF via POST /api/resumes/upload. resumes.py saves unique file to uploads/, invokes scripts/parse_resume.py (PyMuPDF) to extract candidate profile and technical skills, and stores JSON records in resumes table.", "Resume Upload to Ingestion:")
    add_bullet(" User uploads JD via POST /api/job-descriptions/upload. jd_service.py parses PDF, filters technical sections, maps skills against the active resume, and structures priorities into Group A (overlap) and Group B (gaps).", "Job Description Mapping:")
    add_bullet(" When an interview starts, create_interview_session() in interview_service.py initializes AdaptiveState, selects the starting skill and Bloom level 1 (Remember), queries ChromaDB technical_kb collection, prompts Groq API, and generates Question #1.", "Session Setup to Question Generation:")
    add_bullet(" Candidate submits answer to POST /api/interviews/{id}/questions/{q_id}/answer. interview_service.py records answer, invokes evaluation_service.py (Groq critique + SBERT cosine similarity), writes answer_evaluations row, updates AdaptiveState, and generates Question #2.", "Answer Submission to Adaptive Loop:")

    # ============================================================
    # SECTION 3
    # ============================================================
    add_heading_styled("3. Complete Folder Structure", 1)
    tree_text = (
        "SmartInterview/\n"
        "├── .env                             # Active root environment configuration\n"
        "├── .env.example                     # Environment configuration template\n"
        "├── requirements.txt                 # Root Python dependencies\n"
        "├── start_backend.ps1                # PowerShell startup script for FastAPI (Port 8000)\n"
        "├── start_frontend.ps1               # PowerShell startup script for Vite (Port 5174)\n"
        "├── chroma_db/                       # ChromaDB persistent vector database storage\n"
        "│   └── chroma.sqlite3               # Vector database index and SQLite metadata\n"
        "├── data/                            # Knowledge base data assets\n"
        "│   ├── knowledge_stats.json         # Exact audit of 402 chunks, 50 concepts, 8 domains\n"
        "│   ├── raw/                         # Raw scraped technical articles (HTML / text)\n"
        "│   ├── processed/                   # 50 JSON chunk files organized across 8 domain folders\n"
        "│   │   ├── cn/                      # Computer Networking (5 concepts, 33 chunks)\n"
        "│   │   ├── dbms/                    # Database Management Systems (6 concepts, 45 chunks)\n"
        "│   │   ├── design-patterns/         # Software Design Patterns (4 concepts, 105 chunks)\n"
        "│   │   ├── dsa/                     # Data Structures & Algorithms (15 concepts, 63 chunks)\n"
        "│   │   ├── ml-dl/                   # Machine Learning & Deep Learning (5 concepts, 35 chunks)\n"
        "│   │   ├── oop/                     # Object-Oriented Programming (5 concepts, 33 chunks)\n"
        "│   │   ├── os/                      # Operating Systems (5 concepts, 35 chunks)\n"
        "│   │   └── system-design/           # Distributed System Design (5 concepts, 53 chunks)\n"
        "│   ├── evaluation/                  # Ground-truth retrieval queries for RAG validation\n"
        "│   └── resumes/                     # Sample PDF resumes for evaluation\n"
        "├── docs/                            # Development milestone documentation (Weeks 5-10)\n"
        "├── prompts/                         # System and user prompt libraries\n"
        "│   ├── question_generation.json     # Prompts for question types, difficulties, dedup\n"
        "│   └── evaluation_prompts.json      # Structured JSON prompt for answer evaluation\n"
        "├── uploads/                         # Stored uploaded user documents\n"
        "│   └── syllabi/                     # Stored syllabus files for Syllabus Mode RAG\n"
        "├── scripts/                         # Core ingestion and batch evaluation scripts\n"
        "│   ├── clean_text.py                # Text normalization and boilerplate removal\n"
        "│   ├── chunk_text.py                # Recursive character splitter and token validator\n"
        "│   ├── ingest_vector_db.py          # Embedding generation & ChromaDB upsert pipeline\n"
        "│   ├── parse_resume.py              # PDF resume parser with generic header cleanup\n"
        "│   ├── generate_question.py         # Skill taxonomy, RAG retrieval & Groq generator\n"
        "│   ├── test_retrieval.py            # CLI tool to test vector retrieval with domain filters\n"
        "│   ├── evaluate_retrieval.py        # Hit@k and MRR metric benchmark script\n"
        "│   └── validate_chunks.py           # Verification script for chunk length and quality\n"
        "├── backend/                         # FastAPI Backend Application\n"
        "│   ├── requirements.txt             # Backend-specific Python requirements\n"
        "│   ├── setup_database.sql           # Base MySQL database schema (Week 7)\n"
        "│   ├── setup_database_v2.sql        # Migration v2 (Bloom & evaluation tables)\n"
        "│   ├── setup_database_v3.sql        # Migration v3 (Completion reason column)\n"
        "│   ├── setup_database_v4.sql        # Migration v4 (Job Description schema)\n"
        "│   ├── setup_database_v5.sql        # Migration v5 (Candidate Name schema)\n"
        "│   └── app/                         # Application package\n"
        "│       ├── config.py                # Pydantic/dotenv settings and project paths\n"
        "│       ├── database.py              # SQLAlchemy engine, SessionLocal & get_db dependency\n"
        "│       ├── main.py                  # FastAPI initialization, CORS middleware, router mount\n"
        "│       ├── dependencies/auth.py     # HTTPBearer JWT validation dependency (get_current_user)\n"
        "│       ├── models/                  # SQLAlchemy ORM models\n"
        "│       │   ├── user.py              # User model (id, name, email, password_hash)\n"
        "│       │   ├── resume.py            # Resume model (filename, skills JSON, raw_text)\n"
        "│       │   ├── job_description.py   # JobDescription model (skills JSON, raw_text)\n"
        "│       │   ├── interview.py         # InterviewSession model (adaptive_state JSON)\n"
        "│       │   └── question.py          # InterviewQuestion, Answer, AnswerEvaluation models\n"
        "│       ├── schemas/                 # Pydantic request and response schemas\n"
        "│       ├── routers/                 # FastAPI route controllers\n"
        "│       │   ├── auth.py              # POST /register, POST /login, GET /me\n"
        "│       │   ├── resumes.py           # POST /upload, GET /current, DELETE /{id}\n"
        "│       │   ├── job_descriptions.py  # POST /upload, GET /current, GET /mapping\n"
        "│       │   ├── interviews.py        # POST /start, POST /questions/{id}/answer, GET /results\n"
        "│       │   └── users.py             # GET /profile, PUT /profile, GET /performance\n"
        "│       └── services/                # Core business logic services\n"
        "│           ├── auth_service.py      # bcrypt password verification and JWT issuance\n"
        "│           ├── resume_service.py    # Wrapper executing scripts/parse_resume.py\n"
        "│           ├── jd_service.py        # Deterministic JD parser and skill mapping engine\n"
        "│           ├── question_service.py  # Singleton RAG and Groq manager, question generation\n"
        "│           ├── bloom.py             # Bloom's Taxonomy 6-level progression definitions\n"
        "│           ├── adaptive_engine.py   # Deterministic progression rules and skill rotation\n"
        "│           ├── evaluation_service.py# 5-factor scoring engine and LLM evaluator\n"
        "│           ├── interview_service.py # Orchestrator for session lifecycle and transitions\n"
        "│           ├── syllabus_rag_service.py # Dynamic temporary RAG for coursework materials\n"
        "│           └── syllabus_engine.py   # Topic selection with syllabus coverage constraints\n"
        "└── frontend/                        # React 18 / Vite Client Application\n"
        "    ├── package.json                 # Node dependencies (Axios, React Router, Tailwind)\n"
        "    ├── vite.config.js               # Dev server configuration (Port 5174, /api proxy)\n"
        "    └── src/                         # Source directory\n"
        "        ├── main.jsx                 # React root DOM mounting\n"
        "        ├── App.jsx                  # Route definitions and AppLayout wrapper\n"
        "        ├── index.css                # Tailwind CSS v4 styling rules\n"
        "        ├── context/AuthContext.jsx  # Global auth state, login/logout, /auth/me refresh\n"
        "        ├── services/api.js          # Configured Axios instance with auth interceptors\n"
        "        ├── components/              # Reusable layout and navigation components\n"
        "        │   ├── Navbar.jsx           # Top navigation bar for public routes\n"
        "        │   ├── Sidebar.jsx          # Collapsible authenticated navigation bar\n"
        "        │   └── ProtectedRoute.jsx   # Route guard redirecting unauthenticated users\n"
        "        └── pages/                   # Application view pages\n"
        "            ├── Landing.jsx          # Public marketing page\n"
        "            ├── Login.jsx            # User login page\n"
        "            ├── Register.jsx         # User registration page\n"
        "            ├── Dashboard.jsx        # Main user portal and preparation metrics\n"
        "            ├── ResumeUpload.jsx     # Dual resume and JD upload & skill mapping UI\n"
        "            ├── SyllabusUpload.jsx   # Multi-file course syllabus ingestion UI\n"
        "            ├── InterviewSetup.jsx   # Interactive skill focus selection UI\n"
        "            ├── Interview.jsx        # Real-time adaptive questioning interface\n"
        "            ├── Results.jsx          # In-depth post-interview performance analysis\n"
        "            ├── History.jsx          # Chronological record of completed sessions\n"
        "            ├── Profile.jsx          # User details and account management\n"
        "            └── PerformanceDashboard.jsx # Cross-interview aggregated analytics"
    )
    add_code_block(tree_text)

    # File Inventory Table
    file_table = [
        ["File", "Purpose", "Primary Imports / Used By", "Runtime Role"],
        ["backend/app/main.py", "FastAPI entrypoint and CORS configuration", "Imports routers, starts FastAPI app", "Runtime API Entry"],
        ["backend/app/config.py", "Global configuration and path management", "Loads .env, imported by all services", "Configuration"],
        ["backend/app/database.py", "SQLAlchemy database connection setup", "Creates engine, sessionmaker, get_db", "Database Access"],
        ["backend/app/dependencies/auth.py", "JWT token validation dependency", "HTTPBearer, jose.jwt, User model", "Security / Auth"],
        ["backend/app/services/interview_service.py", "Orchestrates full interview lifecycle", "Interviews router, question/eval services", "Core Orchestrator"],
        ["backend/app/services/adaptive_engine.py", "Deterministic adaptive progression logic", "interview_service.py, bloom.py", "Core Algorithm"],
        ["backend/app/services/bloom.py", "Defines 6 Bloom cognitive levels", "adaptive_engine.py, question_service.py", "Domain Model"],
        ["backend/app/services/evaluation_service.py", "Multi-factor answer evaluation engine", "interview_service.py, Groq, SBERT", "AI Evaluation"],
        ["backend/app/services/question_service.py", "Singleton manager for RAG & Groq", "interview_service.py, scripts/generate_question", "Question AI"],
        ["backend/app/services/jd_service.py", "Deterministic JD parser & skill mapper", "job_descriptions router, PyMuPDF", "Document AI"],
        ["backend/app/services/syllabus_rag_service.py", "Temporary isolated RAG for course materials", "interviews router, ChromaDB, PyMuPDF", "Coursework AI"],
        ["scripts/generate_question.py", "Canonical skill taxonomy & RAG pipeline", "question_service.py, parse_resume.py", "Central Knowledge Engine"],
        ["scripts/parse_resume.py", "Deterministic PDF resume parser", "resume_service.py, PyMuPDF", "Resume Extractor"],
        ["scripts/ingest_vector_db.py", "Vector ingestion script for ChromaDB", "SentenceTransformer, ChromaDB", "Data Ingestion Script"],
        ["frontend/src/context/AuthContext.jsx", "Global user state & token refresh", "App.jsx, api.js, components", "Frontend State"],
        ["frontend/src/services/api.js", "Axios client with JWT interceptors", "All frontend pages and AuthContext", "API Communication"],
        ["frontend/src/pages/Interview.jsx", "Interactive real-time interview interface", "React, Lucide, api.js", "Main User View"]
    ]
    add_table_data(file_table[0], file_table[1:], [1.8, 2.2, 1.8, 1.2])

    # ============================================================
    # SECTION 4
    # ============================================================
    add_heading_styled("4. Frontend — Complete Explanation", 1)
    add_p("The frontend application is built as a single-page React 18 application bundled with Vite. It features a complete authentication-aware architecture with route guarding, shared navigation layouts, interactive forms, and real-time state feedback.")

    add_heading_styled("Application Startup & Root Mounting", 2)
    add_bullet(" Execution starts at frontend/index.html which defines the root DOM element <div id=\"root\"></div> and loads src/main.jsx.", "Entrypoint:")
    add_bullet(" main.jsx mounts the root React component App inside React.StrictMode.", "Mounting:")
    add_bullet(" App.jsx wraps the application inside BrowserRouter and AuthProvider, establishing global routing and authentication context.", "Context Setup:")

    add_heading_styled("Route Guarding and Layout Structure", 2)
    add_p("All routes in App.jsx are divided into Public and Protected:")
    add_bullet(" Landing (/), Login (/login), Register (/register). Accessible to anyone without a token.", "Public Routes:")
    add_bullet(" Dashboard (/dashboard), Resume (/resume), Syllabus (/syllabus), Setup (/setup), Results (/results/:id), History (/history), Profile (/profile), and Analytics (/performance).", "Protected Routes:")
    add_bullet(" Wrapped inside ProtectedRoute.jsx. If AuthContext indicates loading=true, it renders a loading spinner. If isAuthenticated=false, it issues a redirect to /login using React Router's <Navigate to=\"/login\" replace />.", "Protection Mechanism:")
    add_bullet(" Authenticated pages are wrapped inside AppLayout, which embeds Sidebar.jsx on the left and renders page children in a scrollable main viewport.", "Layout Wrapper:")

    add_heading_styled("Component-by-Component Analysis", 2)
    add_bullet(" Handles user login with email and password fields. Invokes login() in AuthContext which POSTs to /api/auth/login. On success, persists access_token and user object into localStorage and redirects to /dashboard.", "Login.jsx:")
    add_bullet(" Handles candidate registration with name, email, password, and confirm_password fields. Validates client-side matching, invokes register() in AuthContext, saves the returned JWT, and navigates to /dashboard.", "Register.jsx:")
    add_bullet(" Primary portal for returning candidates. Issues parallel API requests to /api/users/profile, /api/resumes/current, and /api/users/stats. Displays user greeting, preparation status, interview history summary, and quick-action buttons.", "Dashboard.jsx:")
    add_bullet(" Dual upload interface. Allows candidates to drag-and-drop or upload their PDF resume (POST /api/resumes/upload) and Job Description PDF (POST /api/job-descriptions/upload). Once both exist, calls GET /api/job-descriptions/mapping to display matched skills (Group A) and gap skills (Group B).", "ResumeUpload.jsx:")
    add_bullet(" Coursework interview preparation. Accepts multiple PDF, TXT, or DOCX files (POST /api/interviews/upload-syllabus). Displays processing status steps while backend infers subjects and topics, then displays selectable topic tags and stage timing metrics.", "SyllabusUpload.jsx:")
    add_bullet(" Fetches active skills via /api/job-descriptions/mapping or /api/resumes/current. Allows candidates to toggle target skills. On click 'Start Interview', POSTs to /api/interviews/start and redirects to /interview/:id.", "InterviewSetup.jsx:")
    add_bullet(" The core interview view. Displays current question number, target skill, difficulty badge, and Bloom's Taxonomy cognitive level badge. Contains an answer textarea and submit button. Submits answer to POST /api/interviews/{id}/questions/{q_id}/answer, renders structured evaluation feedback for 3 seconds, and seamlessly transitions to the next dynamically generated question.", "Interview.jsx:")
    add_bullet(" Detailed post-interview review. Fetches GET /api/interviews/{id}/results. Renders overall percentage score, per-skill progress bars, cognitive depth milestones, expandable accordion cards for every question with candidate answer and AI feedback, and personalized improvement recommendations.", "Results.jsx:")
    add_bullet(" Chronological log of past sessions fetched from GET /api/interviews/history. Displays interview date, difficulty, questions answered, average score, completion reason, and links to review results.", "History.jsx:")
    add_bullet(" Aggregated longitudinal analytics from GET /api/users/performance. Shows average score across all sessions, total questions answered, identified technical strengths, weak areas needing reinforcement, and recommended study topics.", "PerformanceDashboard.jsx:")

    # ============================================================
    # SECTION 5
    # ============================================================
    add_heading_styled("5. Authentication and Multi-User Behavior", 1)
    add_p("SmartInterview implements stateless JWT (JSON Web Token) bearer authentication using symmetric HMAC-SHA256 signing. Password storage uses bcrypt with salt generation.")

    add_heading_styled("Complete Authentication Flow", 2)
    auth_flow = (
        "Candidate                Frontend (React)           FastAPI Backend           MySQL DB\n"
        "    |                           |                          |                      |\n"
        "    |--- Enter Credentials ---->|                          |                      |\n"
        "    |    (email, password)      |--- POST /api/auth/login->|                      |\n"
        "    |                           |                          |--- Query user ------>|\n"
        "    |                           |                          |<-- User record ------|\n"
        "    |                           |                          |                      |\n"
        "    |                           |                          | [bcrypt.checkpw()]   |\n"
        "    |                           |                          | [jwt.encode(sub)]    |\n"
        "    |                           |<-- 200 TokenResponse ----|                      |\n"
        "    |                           |    (access_token, user)  |                      |\n"
        "    |                           |                          |                      |\n"
        "    |                           | [localStorage.setItem]   |                      |\n"
        "    |                           | [setUser(user)]          |                      |\n"
        "    |<-- Redirect /dashboard ---|                          |                      |"
    )
    add_code_block(auth_flow)

    add_heading_styled("Detailed Multi-User and Browser Scenarios", 2)
    add_bullet(" Frontend receives access_token and user object. AuthContext updates user state. axios interceptor in api.js attaches 'Authorization: Bearer <token>' to every outgoing HTTP request. User is routed to /dashboard.", "What happens when User A logs in?:")
    add_bullet(" React state resets. AuthContext executes its initial useEffect, reading 'token' from localStorage. It dispatches a verification request to GET /api/auth/me. If token is valid, User A's identity is restored without re-logging in. If token expired, localStorage is cleared and user is redirected to /login.", "What happens when User A refreshes the page?:")
    add_bullet(" AuthContext removes 'token' and 'user' from localStorage, resets user state to null, and navigates to Landing page (/). Subsequent protected API calls immediately trigger 401 Unauthorized.", "What happens when User A logs out?:")
    add_bullet(" Because localStorage is shared across all browser tabs for the same origin (e.g. http://localhost:5174), logging in as User B in Tab 2 overwrites User A's token in localStorage. Tab 1 will begin transmitting User B's JWT on subsequent API requests. The codebase does NOT implement multi-tab user isolation; this is documented as an architectural characteristic.", "What happens when User B logs in on another tab?:")
    add_bullet(" If an expired or tampered JWT exists in localStorage, GET /api/auth/me fails with 401. AuthContext catches this, purges localStorage, and redirects to /login.", "What happens if an invalid JWT remains in storage?:")
    add_bullet(" The Axios response interceptor in api.js captures response.status === 401, removes 'token' and 'user' from localStorage, and forces window.location.href = '/login'.", "What happens when a protected API returns 401?:")

    # Save initial version to disk to ensure no data loss
    doc.save("SmartInterview_Master_Technical_Documentation.docx")
    print("Checkpoint saved: Sections 1-5.")

    # ============================================================
    # SECTION 6
    # ============================================================
    add_heading_styled("6. Backend — Complete Explanation", 1)
    add_p("The backend is built with FastAPI and runs on Uvicorn. It exposes a clean REST API divided into modular domain routers.")

    endpoints = [
        ["Method", "Endpoint", "Purpose", "Auth Required", "Database / External Services"],
        ["POST", "/api/auth/register", "Register new user account", "No", "users table (INSERT), bcrypt"],
        ["POST", "/api/auth/login", "Authenticate & issue JWT", "No", "users table (SELECT), bcrypt"],
        ["GET", "/api/auth/me", "Verify token & return user", "Yes (Bearer)", "users table (SELECT)"],
        ["POST", "/api/resumes/upload", "Upload & parse PDF resume", "Yes (Bearer)", "resumes table, PyMuPDF, file storage"],
        ["GET", "/api/resumes/current", "Get current active resume", "Yes (Bearer)", "resumes table (SELECT)"],
        ["DELETE", "/api/resumes/{id}", "Delete uploaded resume", "Yes (Bearer)", "resumes table (DELETE), file unlink"],
        ["POST", "/api/job-descriptions/upload", "Upload & parse target JD", "Yes (Bearer)", "job_descriptions table, PyMuPDF"],
        ["GET", "/api/job-descriptions/current", "Get current active JD", "Yes (Bearer)", "job_descriptions table (SELECT)"],
        ["GET", "/api/job-descriptions/mapping", "Resume vs JD skill mapping", "Yes (Bearer)", "resumes + job_descriptions tables"],
        ["POST", "/api/interviews/start", "Start adaptive interview session", "Yes (Bearer)", "interview_sessions, ChromaDB, Groq API"],
        ["GET", "/api/interviews/history", "List past interview sessions", "Yes (Bearer)", "interview_sessions, answers tables"],
        ["GET", "/api/interviews/{id}", "Get session state & questions", "Yes (Bearer)", "interview_sessions, interview_questions"],
        ["POST", "/api/interviews/{id}/questions/{qid}/answer", "Submit answer & get evaluation", "Yes (Bearer)", "answers, answer_evaluations, Groq, SBERT"],
        ["POST", "/api/interviews/{id}/complete", "Mark interview completed", "Yes (Bearer)", "interview_sessions (UPDATE completed_at)"],
        ["GET", "/api/interviews/{id}/results", "Get rich interview results", "Yes (Bearer)", "interview_sessions, evaluations, questions"],
        ["POST", "/api/interviews/upload-syllabus", "Upload course materials (RAG)", "Yes (Bearer)", "ChromaDB temporary collection, PyMuPDF"],
        ["GET", "/api/users/profile", "Get profile details & stats", "Yes (Bearer)", "users, resumes, interview_sessions"],
        ["PUT", "/api/users/profile", "Update profile details", "Yes (Bearer)", "users table (UPDATE)"],
        ["GET", "/api/users/performance", "Aggregated performance analytics", "Yes (Bearer)", "answer_evaluations across all sessions"]
    ]
    add_table_data(endpoints[0], endpoints[1:], [0.8, 2.5, 2.0, 1.0, 1.7])

    # ============================================================
    # SECTION 7
    # ============================================================
    add_heading_styled("7. Database — Complete Explanation", 1)
    add_p("The relational persistence layer uses MySQL 8.0 with InnoDB tables, utf8mb4 encoding, foreign key constraints with CASCADE deletes, and JSON columns for dynamic nested metadata.")

    er_diagram = (
        "+-----------------------------------------------------------------------------------+\n"
        "|                                       users                                       |\n"
        "|  id (PK, AI) | name | email (UNIQUE) | password_hash | created_at | updated_at    |\n"
        "+-----------------------------------------------------------------------------------+\n"
        "        | 1                                       | 1                       | 1\n"
        "        |                                         |                         |\n"
        "        v *                                       v *                       v *\n"
        "+-----------------------------+  +-------------------------------+  +---------------+ \n"
        "|           resumes           |  |        job_descriptions       |  |interview_sess |\n"
        "| id (PK, AI)                 |  | id (PK, AI)                   |  | id (PK, AI)   |\n"
        "| user_id (FK -> users.id)    |  | user_id (FK -> users.id)      |  | user_id (FK)  |\n"
        "| filename | file_path | name |  | filename | file_path          |  | resume_id (FK)|\n"
        "| skills (JSON)               |  | skills (JSON)                 |  | jd_id (FK)    |\n"
        "| projects | exp | edu (JSON) |  | raw_text (MEDIUMTEXT)         |  | difficulty    |\n"
        "| raw_text (MEDIUMTEXT)       |  +-------------------------------+  | question_type |\n"
        "+-----------------------------+                                     | question_count|\n"
        "                                                                    | selected_skill|\n"
        "                                                                    | status        |\n"
        "                                                                    | is_adaptive   |\n"
        "                                                                    | adaptive_state|\n"
        "                                                                    | cur_bloom_lvl |\n"
        "                                                                    | final_recomms |\n"
        "                                                                    | comp_reason   |\n"
        "                                                                    | mode / syll_id|\n"
        "                                                                    +---------------+\n"
        "                                                                            |\n"
        "                                                                            v 1..*\n"
        "                                                                    +---------------+\n"
        "                                                                    |interview_quest|\n"
        "                                                                    | id (PK, AI)   |\n"
        "                                                                    | session_id(FK)|\n"
        "                                                                    | question_num  |\n"
        "                                                                    | skill         |\n"
        "                                                                    | question_type |\n"
        "                                                                    | difficulty    |\n"
        "                                                                    | question_text |\n"
        "                                                                    | rag_context   |\n"
        "                                                                    | project_ctx   |\n"
        "                                                                    | bloom_level   |\n"
        "                                                                    | bloom_lvl_num |\n"
        "                                                                    +---------------+\n"
        "                                                                            |\n"
        "                                                                            v 1..1\n"
        "                                                                    +---------------+\n"
        "                                                                    |    answers    |\n"
        "                                                                    | id (PK, AI)   |\n"
        "                                                                    | question_id FK|\n"
        "                                                                    | session_id FK |\n"
        "                                                                    | user_id FK    |\n"
        "                                                                    | answer_text   |\n"
        "                                                                    | submitted_at  |\n"
        "                                                                    +---------------+\n"
        "                                                                            |\n"
        "                                                                            v 1..1\n"
        "                                                                    +---------------+\n"
        "                                                                    |answer_evaluat.|\n"
        "                                                                    | id (PK, AI)   |\n"
        "                                                                    | answer_id (FK)|\n"
        "                                                                    | question_id(FK|\n"
        "                                                                    | technical_scr |\n"
        "                                                                    | complete_scr  |\n"
        "                                                                    | relevance_scr |\n"
        "                                                                    | sem_sim_score |\n"
        "                                                                    | concept_cov   |\n"
        "                                                                    | overall_score |\n"
        "                                                                    | feedback(TEXT)|\n"
        "                                                                    | strengths JSON|\n"
        "                                                                    | weakness JSON |\n"
        "                                                                    | expected_c JSON\n"
        "                                                                    | found_c (JSON)|\n"
        "                                                                    +---------------+"
    )
    add_code_block(er_diagram)

    # ============================================================
    # SECTION 8 & 9
    # ============================================================
    add_heading_styled("8. Knowledge Base and RAG — In-Depth Analysis", 1)
    add_p("The SmartInterview Retrieval-Augmented Generation (RAG) system anchors interview questions in verified computer science literature. The baseline knowledge base contains exactly 402 chunks across 50 concepts in 8 technical domains, verified directly from data/knowledge_stats.json.")

    add_heading_styled("Knowledge Base Ingestion Pipeline", 2)
    add_bullet(" High-quality technical articles and documentation were scraped from authoritative sources such as GeeksforGeeks and GitHub technical roadmaps into data/raw/.", "1. Raw Source Collection:")
    add_bullet(" clean_text.py removes web navigation boilerplate, ads, cookie disclaimers, and formatting noise.", "2. Cleaning & Normalization:")
    add_bullet(" chunk_text.py splits clean text using LangChain's RecursiveCharacterTextSplitter (chunk_size=1024, chunk_overlap=200). Chunks are tokenized using tiktoken (cl100k_base). Chunks outside 50-400 tokens are discarded.", "3. Token-Bounded Chunking:")
    add_bullet(" Each valid chunk is formatted into structured JSON with chunk_id, domain, concept, source, source_url, text, and token_count.", "4. Chunk Formatting:")
    add_bullet(" ingest_vector_db.py reads all 50 JSON files in data/processed/, encodes text using SentenceTransformer('all-MiniLM-L6-v2'), and batch-upserts the 402 documents with metadata into ChromaDB collection 'technical_kb'.", "5. Vector DB Upsert:")

    add_heading_styled("Domain Breakdown & Chunk Distribution", 2)
    kb_summary = [
        ["Domain Code", "Full Domain Name", "Concepts Count", "Total Chunks", "Key Covered Concepts"],
        ["cn", "Computer Networks", "5", "33", "dns (8), http-vs-https (5), ipv4-vs-ipv6 (5), osi-model (13), tcp-vs-udp (2)"],
        ["dbms", "Database Systems", "6", "45", "acid-properties (8), indexing (9), keys (10), normalization (7), sql-joins (6), transactions (5)"],
        ["design-patterns", "Software Patterns", "4", "105", "factory (10), observer (34), singleton (26), strategy (35)"],
        ["dsa", "Data Structures & Algos", "15", "63", "array (9), graph (8), linked-list (7), string (6), tree (6), matrix (4), recursion (4), hash-table (3)"],
        ["ml-dl", "Machine Learning & DL", "5", "35", "decision-trees (11), neural-networks (9), bias-variance (5), linear-vs-logistic (5), overfitting (5)"],
        ["oop", "Object-Oriented Prog.", "5", "33", "polymorphism (10), composition-inheritance (7), abstraction (6), interfaces (6), encapsulation (4)"],
        ["os", "Operating Systems", "5", "35", "scheduling-algorithms (10), virtual-memory (10), mutex-vs-semaphore (8), process-vs-thread (4), deadlock (3)"],
        ["system-design", "Distributed Systems", "5", "53", "consistent-hashing (16), load-balancing (13), caching (11), cap-theorem (9), microservices (4)"],
        ["TOTAL", "8 Domains", "50 Concepts", "402 Chunks", "100% PASS Quality Audit (knowledge_stats.json)"]
    ]
    add_table_data(kb_summary[0], kb_summary[1:], [1.2, 1.8, 1.0, 1.0, 3.0])

    add_heading_styled("9. Knowledge Base Data Structure", 1)
    add_p("Every chunk stored in data/processed/ adheres strictly to the following JSON schema:")
    chunk_json_sample = (
        "{\n"
        "  \"chunk_id\": \"dbms_acid-properties_001\",\n"
        "  \"domain\": \"dbms\",\n"
        "  \"concept\": \"acid-properties\",\n"
        "  \"source\": \"geeksforgeeks\",\n"
        "  \"source_url\": \"https://www.geeksforgeeks.org/acid-properties-in-dbms/\",\n"
        "  \"text\": \"Transactions are fundamental operations that allow us to modify and retrieve data. However, to ensure...\",\n"
        "  \"token_count\": 179\n"
        "}"
    )
    add_code_block(chunk_json_sample)

    # ============================================================
    # SECTION 10
    # ============================================================
    add_heading_styled("10. Resume Processing", 1)
    add_p("Resume processing is implemented deterministically without LLM calls in scripts/parse_resume.py and wrapped by backend/app/services/resume_service.py.")
    
    add_bullet(" PyMuPDF (pymupdf.open) extracts raw text across all pages. clean_document_pages() inspects header and footer lines across pages, identifying and purging repeated header/footers, pagination markers ('Page 1 of 2', '- 2 -'), and continuation headers.", "Text Extraction & Cleanup:")
    add_bullet(" Scans top 10 lines of document text, filtering out contact lines (@, http, phone numbers) and reject keywords ('Resume', 'Curriculum Vitae', 'Developer'), validating against a Title-Case regex.", "Candidate Name Extraction:")
    add_bullet(" Regex scans for common section headings (Skills, Projects, Experience, Education, Certifications) and partitions text into distinct section blocks.", "Section Boundary Detection:")
    add_bullet(" Central matching engine from scripts/generate_question.py matches canonical skills using phrase-length descending order, regex word boundaries, and overlap suppression (preventing 'Java' from matching inside 'JavaScript').", "Canonical Skill Matching:")

    # ============================================================
    # SECTION 11 & 12
    # ============================================================
    add_heading_styled("11. Question Generation Pipeline", 1)
    add_p("Question generation executes dynamically on-demand for each interview step:")
    add_bullet(" Candidate skill (e.g. 'PostgreSQL'), current difficulty (e.g. 'medium'), current Bloom level (e.g. 'Apply'), candidate project context (optional), and previously asked questions for deduplication.", "Generator Inputs:")
    add_bullet(" SKILL_DOMAIN_MAP resolves 'postgresql' to ['dbms']. DOMAIN_QUERY_TEMPLATES builds domain-aware query: 'SQL database postgresql concepts: indexing, joins, normalization, transactions, keys, ACID'.", "Domain-Aware Query Construction:")
    add_bullet(" retrieve_rag_context() encodes query using SentenceTransformer and queries ChromaDB with where={'domain': 'dbms'}, retrieving top 3 chunks.", "Domain-Filtered Retrieval:")
    add_bullet(" Combines system prompt from prompts/question_generation.json, difficulty instruction, question type guidance, retrieved knowledge chunks, and previous questions list.", "Prompt Construction:")
    add_bullet(" call_groq() submits prompt to Groq API (openai/gpt-oss-120b) with temperature 0.7. If Groq truncates or errors, exponential backoff retries (5s, 10s, 20s) with increased token budget.", "LLM Call & Retry Logic:")

    add_heading_styled("12. Answer Evaluation Engine", 1)
    add_p("The evaluation service (backend/app/services/evaluation_service.py) scores candidate responses using a hybrid framework:")
    add_code_block("Overall Score = (0.30 * Technical) + (0.20 * Completeness) + (0.20 * Relevance) + (0.15 * Semantic Similarity) + (0.15 * Concept Coverage)")
    
    add_bullet(" Evaluates technical correctness, depth, relevance, and identifies specific strengths, weaknesses, expected concepts, and found concepts. Uses temperature 0.3 for consistency.", "LLM Evaluation (Groq):")
    add_bullet(" Encodes candidate answer and reference text (question + RAG chunks up to 1500 chars) using SentenceTransformer. Calculates cosine similarity and applies non-linear scaling: max(0, min(1, (cosine - 0.1) / 0.6)) * 100. Measures semantic proximity, not correctness.", "Semantic Similarity (SBERT):")
    add_bullet(" Averages the LLM-reported concept score with algorithmic calculation: round(100 * len(found_concepts) / len(expected_concepts)).", "Concept Coverage:")

    # ============================================================
    # SECTION 13 & 14
    # ============================================================
    add_heading_styled("13. Adaptive Learning Engine", 1)
    add_p("The adaptive learning engine (backend/app/services/adaptive_engine.py) governs interview difficulty and question selection deterministically. The LLM does NOT decide interview flow.")

    add_heading_styled("Deterministic Score-Based Progression Rules", 2)
    rules_table = [
        ["Evaluation Score", "Bloom's Level Action", "Difficulty Action", "System Explanation"],
        ["Score >= 80", "Advance to Next Bloom Level", "If at Create (Max), Advance Difficulty", "Candidate showed mastery; increase cognitive complexity."],
        ["Score 50 - 79", "Maintain Current Bloom Level", "Maintain Current Difficulty", "Solid performance; reinforce current level with new topic."],
        ["Score < 50", "Regress to Previous Bloom Level", "If at Remember (Min), Lower Difficulty", "Candidate struggled; decrease complexity to foundational level."]
    ]
    add_table_data(rules_table[0], rules_table[1:], [1.5, 2.0, 2.0, 2.5])

    add_heading_styled("Skill Selection Algorithm", 2)
    add_p("select_skill() executes a weakness-biased round-robin algorithm:")
    add_bullet("Calculates total attempts and average score for all selected skills.", level=0)
    add_bullet("Finds the minimum number of attempts across the skill pool.", level=0)
    add_bullet("Filters candidates to only those skills with the minimum attempts (ensures coverage).", level=0)
    add_bullet("If multiple skills tie for minimum attempts, selects the skill with the lowest average score (prioritizes weakness).", level=0)

    add_heading_styled("Interview Completion & Termination Safeguards", 2)
    add_bullet(" Candidates can click 'Finish Interview' at any time. complete_interview() records completion_reason='manual' and computes final recommendations.", "Standard Completion (Manual):")
    add_bullet(" Triggers ONLY when all 4 conditions are met: (1) At least 5 questions answered, (2) Current difficulty is 'easy', (3) Last 3 consecutive answers all scored < 40%, (4) Clear evidence candidate cannot handle fundamentals. completion_reason='auto_stop_fundamental_struggle'.", "Strict Auto-Stop Safeguard:")
    add_bullet(" Hard boundary of MAX_ADAPTIVE_QUESTIONS = 30 to prevent infinite server loops. completion_reason='max_questions_safety_limit'.", "Technical Safety Limit:")

    add_heading_styled("14. Bloom's Taxonomy Representation", 1)
    add_p("SmartInterview models Bloom's Revised Taxonomy across 6 immutable levels defined in backend/app/services/bloom.py:")
    bloom_table = [
        ["Order", "Level Name", "Cognitive Focus", "Preferred Question Types"],
        ["1", "Remember", "Recall facts, terminology, definitions", "conceptual, practical, technical_reasoning"],
        ["2", "Understand", "Explain concepts, summarize, interpret meaning", "conceptual, technical_reasoning, practical"],
        ["3", "Apply", "Use knowledge in practical code/problem solving", "practical, conceptual, scenario"],
        ["4", "Analyze", "Compare approaches, break down trade-offs", "technical_reasoning, practical, scenario"],
        ["5", "Evaluate", "Judge solutions, justify architectural decisions", "scenario, technical_reasoning, project"],
        ["6", "Create", "Architect new systems, synthesize designs", "project, scenario, practical"]
    ]
    add_table_data(bloom_table[0], bloom_table[1:], [0.8, 1.4, 2.8, 2.5])

    doc.save("SmartInterview_Master_Technical_Documentation.docx")
    print("Checkpoint saved: Sections 6-14.")

    # ============================================================
    # SECTION 15 & 16
    # ============================================================
    add_heading_styled("15. Complete Interview Lifecycle", 1)
    add_p("A complete interview session progresses through 12 distinct stages:")
    add_bullet("Candidate logs in, views Dashboard, and navigates to Resume Upload.", "Stage 1 — Access:")
    add_bullet("Candidate uploads resume PDF and job description PDF. Deterministic matching extracts skills and identifies target focus areas.", "Stage 2 — Ingestion:")
    add_bullet("Candidate selects target skills on Interview Setup and clicks Start.", "Stage 3 — Configuration:")
    add_bullet("FastAPI backend creates InterviewSession row with is_adaptive=True and initialized AdaptiveState.", "Stage 4 — Session Creation:")
    add_bullet("Adaptive engine selects first skill, Bloom Level 1 (Remember), queries ChromaDB, and prompts Groq LLM to generate Question #1.", "Stage 5 — Q1 Generation:")
    add_bullet("Frontend receives and displays Question #1. Candidate enters text response and clicks Submit.", "Stage 6 — Answering:")
    add_bullet("Backend evaluates response using Groq critique, SBERT semantic similarity, and concept coverage math, writing an AnswerEvaluation row.", "Stage 7 — Evaluation:")
    add_bullet("Evaluation modal displays score and feedback on candidate's screen for 3 seconds.", "Stage 8 — Feedback:")
    add_bullet("Adaptive engine updates LearnerProfile for that skill. If score >= 80, Bloom advances to Understand. Skill selection rotates to least-attempted/weakest skill.", "Stage 9 — Adaptation:")
    add_bullet("Process loops for subsequent questions, adapting difficulty and cognitive depth dynamically.", "Stage 10 — Loop:")
    add_bullet("Candidate clicks 'Finish Interview'. Session status is updated to 'completed' with completion_reason='manual'.", "Stage 11 — Completion:")
    add_bullet("Candidate is routed to Results page displaying overall score, per-skill mastery, cognitive milestones, question breakdowns, and AI recommendations.", "Stage 12 — Results:")

    add_heading_styled("16. End-to-End Data Flow", 1)
    add_p("Data movement across major actions traces directly through the code:")
    add_bullet("Payload: {email, password} -> POST /api/auth/login -> auth_service.authenticate_user() -> bcrypt check -> create_access_token() -> Response: {access_token, user} -> localStorage.", "Login Flow:")
    add_bullet("FormData: {file: PDF} -> POST /api/resumes/upload -> write to uploads/ -> parse_resume.py (PyMuPDF) -> extract skills & projects -> INSERT into resumes -> Response: ResumeResponse JSON.", "Resume Upload Flow:")
    add_bullet("Payload: {selected_skills} -> POST /api/interviews/start -> create_interview_session() -> initialize_state() -> generate_single_question() -> ChromaDB query -> Groq call -> INSERT session & question #1 -> Response: {session_id, current_question}.", "Interview Start Flow:")
    add_bullet("Payload: {answer_text} -> POST /api/interviews/{id}/questions/{qid}/answer -> submit_and_evaluate() -> INSERT answer -> evaluate_answer() -> INSERT answer_evaluations -> decide_next() -> generate_single_question() -> INSERT question #2 -> Response: {evaluation, next_question, is_complete}.", "Answer & Evaluation Flow:")
    add_bullet("POST /api/interviews/{id}/complete -> complete_interview() -> calculate_session_summary() -> generate_recommendations() -> UPDATE interview_sessions -> Response: SessionResultsResponse.", "Finish Interview Flow:")

    # ============================================================
    # SECTION 17 & 18
    # ============================================================
    add_heading_styled("17. Code-Level Deep Dive", 1)
    add_bullet("backend/app/services/interview_service.py: Contains submit_and_evaluate(), create_interview_session(), complete_interview(), and get_session_results(). Orchestrates the interaction between database models, question generator, evaluator, and adaptive state machine.", "interview_service.py:")
    add_bullet("backend/app/services/adaptive_engine.py: Defines AdaptiveState, SkillPerformance, decide_next(), and select_skill(). Implements deterministic rules: Score >= 80 advances Bloom; Score < 50 regresses Bloom; select_skill() prioritizes minimum attempts then lowest average score.", "adaptive_engine.py:")
    add_bullet("backend/app/services/evaluation_service.py: Defines evaluate_answer(), _llm_evaluate(), and _compute_semantic_similarity(). Implements 30/20/20/15/15 scoring formula, clamps all scores to [0, 100], and protects against JSON parsing errors.", "evaluation_service.py:")
    add_bullet("backend/app/services/question_service.py: Manages lazy-initialized thread-safe singletons for SentenceTransformer, ChromaDB collection, and Groq client. Wraps scripts/generate_question.py without modifying original CLI code.", "question_service.py:")
    add_bullet("backend/app/services/jd_service.py: Extracts technical text from Job Descriptions, filters non-technical sections (salary, benefits), and maps skills into Group A (overlap) and Group B (gaps).", "jd_service.py:")

    add_heading_styled("18. Dependency Map", 1)
    dep_text = (
        "frontend/src/App.jsx\n"
        "  └── AuthContext.jsx -> services/api.js -> Axios -> Backend REST API\n"
        "backend/app/main.py\n"
        "  ├── routers/auth.py -> services/auth_service.py -> models/user.py\n"
        "  ├── routers/resumes.py -> services/resume_service.py -> scripts/parse_resume.py -> PyMuPDF\n"
        "  ├── routers/job_descriptions.py -> services/jd_service.py -> PyMuPDF & scripts/generate_question.py\n"
        "  └── routers/interviews.py -> services/interview_service.py\n"
        "        ├── services/adaptive_engine.py -> services/bloom.py\n"
        "        ├── services/evaluation_service.py -> Groq API & SentenceTransformers\n"
        "        └── services/question_service.py -> ChromaDB & Groq API & scripts/generate_question.py"
    )
    add_code_block(dep_text)

    # ============================================================
    # SECTION 19 & 20
    # ============================================================
    add_heading_styled("19. Configuration and Environment Variables", 1)
    env_table = [
        ["Variable", "Location", "Purpose", "Default Value if Unset", "Sensitivity"],
        ["DATABASE_URL", "Backend (.env)", "SQLAlchemy MySQL connection string", "mysql+pymysql://root:root@localhost:3306/smartinterview", "Secret"],
        ["JWT_SECRET_KEY", "Backend (.env)", "HMAC signing secret for authentication tokens", "dev-secret-change-in-production", "Secret"],
        ["JWT_ALGORITHM", "Backend (.env)", "Cryptographic algorithm for JWT", "HS256", "Public"],
        ["JWT_EXPIRATION_MINUTES", "Backend (.env)", "Token validity lifespan in minutes", "1440 (24 hours)", "Public"],
        ["GROQ_API_KEY", "Backend (.env)", "API key for Groq Cloud LLM inference", "None (Must be provided)", "Secret"],
        ["GROQ_MODEL", "Backend (.env)", "Target LLM model identifier on Groq", "openai/gpt-oss-120b", "Public"]
    ]
    add_table_data(env_table[0], env_table[1:], [1.8, 1.2, 2.2, 1.8, 0.8])

    add_heading_styled("20. Error Handling & Failure Recovery", 1)
    add_bullet("auth.py catches invalid credentials or duplicate emails and returns 401 Unauthorized or 400 Bad Request. Frontend catches error and displays red error banner.", "Authentication Failures:")
    add_bullet("Only PDF files up to 5MB are accepted. If text cannot be extracted or no skills match, uploaded file is deleted and 422 Unprocessable Entity is returned.", "Document Upload Failures:")
    add_bullet("SQLAlchemy uses pool_pre_ping=True and pool_recycle=3600 to transparently reconnect if MySQL drops idle connections.", "Database Disconnections:")
    add_bullet("call_groq() catches network and rate limit errors and applies exponential retry backoff (5s, 10s, 20s). If all retries fail, question text is set to '[GENERATION FAILED]' without crashing the backend.", "Groq API Failures:")
    add_bullet("Wrapped in try-except blocks. If JSON parsing fails or markdown is returned, _safe_defaults() returns fallback score of 0 with constructive error message.", "Evaluation Parsing Errors:")

    # ============================================================
    # SECTION 21 & 22
    # ============================================================
    add_heading_styled("21. Security Posture", 1)
    add_bullet("Passwords hashed with bcrypt, truncated safely to 72 bytes. Plaintext passwords never stored or logged.", "Password Security:")
    add_bullet("Tokens signed with HS256 secret. All protected endpoints verify signature via get_current_user dependency.", "Token Authentication:")
    add_bullet("All database access uses SQLAlchemy ORM parameterized queries; raw SQL concatenation is not used.", "SQL Injection Defense:")
    add_bullet("Uploaded files are stored outside the public document root with sanitized UUID filenames; extension and 5MB size limits enforced.", "File Upload Security:")
    add_bullet("Configured in main.py to allow only frontend origin (localhost:5173 / localhost:5174).", "CORS Configuration:")
    add_bullet("Tokens stored in localStorage are vulnerable to XSS; no refresh token rotation is implemented; multi-tab session isolation is absent.", "Identified Limitations:")

    add_heading_styled("22. Current Implementation vs Planned Features", 1)
    feature_table = [
        ["Feature Area", "Implementation Status", "Evidence in Codebase"],
        ["JWT Authentication", "Implemented", "backend/app/routers/auth.py, dependencies/auth.py, AuthContext.jsx"],
        ["PDF Resume Parsing", "Implemented", "scripts/parse_resume.py, backend/app/services/resume_service.py"],
        ["Job Description Mapping", "Implemented", "backend/app/services/jd_service.py, routers/job_descriptions.py"],
        ["ChromaDB Vector RAG", "Implemented", "scripts/ingest_vector_db.py, backend/app/services/question_service.py"],
        ["Bloom's Taxonomy Progression", "Implemented", "backend/app/services/bloom.py, adaptive_engine.py"],
        ["Deterministic Adaptive Engine", "Implemented", "backend/app/services/adaptive_engine.py (decide_next, select_skill)"],
        ["Multi-Factor Answer Evaluation", "Implemented", "backend/app/services/evaluation_service.py (30/20/20/15/15 formula)"],
        ["Syllabus Mode Coursework RAG", "Implemented", "backend/app/services/syllabus_rag_service.py, syllabus_engine.py"],
        ["Performance Analytics Dashboard", "Implemented", "frontend/src/pages/PerformanceDashboard.jsx, routers/users.py"],
        ["Strict Auto-Stop Safeguard", "Implemented", "backend/app/services/interview_service.py (lines 405-416)"],
        ["Speech-to-Text (Voice Answering)", "Not Implemented / Planned", "Comment in evaluation_service.py: 'deferred to Week 11'"],
        ["Text-to-Speech (Spoken Question)", "Not Implemented / Planned", "No TTS library or audio playback in Interview.jsx"],
        ["Facial / Video Emotion Analysis", "Not Implemented", "No webcam capture or computer vision dependencies present"],
        ["Interview Question Timer", "Not Implemented", "No countdown timer or auto-submission on timeout in Interview.jsx"]
    ]
    add_table_data(feature_table[0], feature_table[1:], [2.2, 1.8, 3.2])

    doc.save("SmartInterview_Master_Technical_Documentation.docx")
    print("Checkpoint saved: Sections 15-22.")

    # ============================================================
    # SECTION 23, 24, 25, 26, 27
    # ============================================================
    add_heading_styled("23. Actual Current Knowledge Base State", 1)
    add_p("The current production baseline knowledge base contains exactly 402 chunks across 50 concepts in 8 technical domains. This was verified through data/knowledge_stats.json. Earlier experimental branches exploring 4,817 chunks are not part of this baseline.")

    add_heading_styled("24. What Happens When the Application Starts?", 1)
    add_bullet("User launches MySQL 8 server on localhost:3306 with database smartinterview.", "1. Database:")
    add_bullet("Executed via start_backend.ps1. Loads .env, starts Uvicorn on port 8000, connects SQLAlchemy connection pool, pre-loads SentenceTransformer model, and connects to ChromaDB at chroma_db/.", "2. Backend Start:")
    add_bullet("Executed via start_frontend.ps1. Runs npm run dev in frontend/, binding Vite dev server to port 5174 and reverse-proxying /api to http://localhost:8000.", "3. Frontend Start:")

    add_heading_styled("25. Complete User Journey: Raj's Experience", 1)
    add_p("Suppose Raj, a software engineering candidate, prepares for a Backend Developer interview:")
    add_bullet("Raj opens http://localhost:5174, clicks 'Get Started', and registers with his email. The backend creates his account and stores a JWT in his browser.", "1. Registration:")
    add_bullet("Raj uploads his resume PDF. PyMuPDF extracts his skills (Python, PostgreSQL, Docker, FastApi) and projects. He then uploads a Backend Job Description, and the system identifies PostgreSQL and FastAPI as high-priority focus skills.", "2. Document Ingestion:")
    add_bullet("Raj navigates to Interview Setup, verifies target skills, and clicks 'Start Interview'. The backend initializes his session at Bloom Level 1 (Remember).", "3. Interview Start:")
    add_bullet("Question #1 appears: 'What are ACID properties in relational databases and what does Atomicity ensure?'. Raj types his answer explaining commit and rollback mechanics.", "4. Question 1:")
    add_bullet("Raj clicks Submit. The system runs Groq evaluation and SBERT similarity, awarding him 88%. The UI shows green score feedback for 3 seconds.", "5. Evaluation:")
    add_bullet("Because score >= 80, the adaptive engine promotes Bloom level to Level 2 (Understand). It rotates to FastAPI and generates Question #2 testing request validation and Pydantic models.", "6. Adaptation:")
    add_bullet("After answering 5 questions across various skills and levels, Raj clicks 'Finish Interview'. Results page displays his 84% average, skill mastery bars, and personalized AI recommendations.", "7. Completion:")

    add_heading_styled("26. Interview Engine Decision Simulations", 1)
    add_bullet("Initial Bloom: Remember (1), Difficulty: Medium. Candidate scores 85%. Rule: Score >= 80 advances Bloom level. Decision: Next Bloom = Understand (2), Difficulty = Medium. Skill selection rotates to next least-attempted skill.", "Simulation 1 (High Score at Remember):")
    add_bullet("Initial Bloom: Create (6), Difficulty: Medium. Candidate scores 92%. Rule: Score >= 80 attempts Bloom advancement. Already at maximum level (Create). Decision: Bloom remains Create (6), Difficulty advances to Hard.", "Simulation 2 (High Score at Max Bloom):")
    add_bullet("Initial Bloom: Apply (3), Difficulty: Medium. Candidate scores 68%. Rule: Score 50-79 maintains level. Decision: Next Bloom = Apply (3), Difficulty = Medium.", "Simulation 3 (Moderate Score):")
    add_bullet("Initial Bloom: Apply (3), Difficulty: Medium. Candidate scores 38%. Rule: Score < 50 regresses level. Decision: Next Bloom = Understand (2), Difficulty = Medium.", "Simulation 4 (Poor Score):")
    add_bullet("Initial Bloom: Remember (1), Difficulty: Medium. Candidate scores 25%. Rule: Score < 50 attempts Bloom regression. Already at minimum level (Remember). Decision: Bloom remains Remember (1), Difficulty drops to Easy.", "Simulation 5 (Poor Score at Min Bloom):")
    add_bullet("Candidate has answered 5 questions, current difficulty is Easy, and the last 3 consecutive questions all scored < 40%. Decision: is_auto_stop triggers True, session automatically completes with completion_reason='auto_stop_fundamental_struggle'.", "Simulation 6 (Strict Auto-Stop Safeguard):")

    add_heading_styled("27. RAG Retrieval & Prompt Construction Examples", 1)
    add_bullet("Skill: SQL -> Domain: dbms -> Query: 'SQL database SQL concepts: indexing, joins, normalization, transactions, keys, ACID' -> Retrieves: dbms_indexing_001, dbms_indexing_002 -> LLM Prompt: Generates B-Tree index lookup question.", "Example 1 (DBMS Indexing):")
    add_bullet("Skill: Java -> Domain: oop -> Query: 'Object-oriented programming Java concepts: classes, inheritance, polymorphism, encapsulation, abstraction' -> Retrieves: oop_polymorphism_001 -> Generates method overriding vs overloading question.", "Example 2 (OOP Polymorphism):")
    add_bullet("Skill: Docker -> Domain: os -> Query: 'Operating systems Docker: processes, threads, memory management, scheduling, synchronization' -> Retrieves: os_process-vs-thread_001 -> Generates container isolation vs virtual machine question.", "Example 3 (OS Process Isolation):")
    add_bullet("Skill: Kafka -> Domain: system-design -> Query: 'System design Kafka: scalability, caching, load balancing, microservices, distributed systems' -> Retrieves: system-design_caching_001, load-balancing_001 -> Generates message queue durability question.", "Example 4 (System Design Caching):")
    add_bullet("Skill: PyTorch -> Domain: ml-dl -> Query: 'Machine learning PyTorch: regression, classification, neural networks, overfitting, bias-variance' -> Retrieves: ml-dl_overfitting-vs-underfitting_001 -> Generates dropout and regularization question.", "Example 5 (ML Overfitting):")

    # ============================================================
    # SECTION 28, 29, 30, 31, 32 & EXECUTIVE SUMMARY
    # ============================================================
    add_heading_styled("28. File-by-File Summary Inventory", 1)
    file_inventory = [
        ["#", "File Path", "Description & Role", "Criticality"],
        ["1", "backend/app/main.py", "FastAPI app instance, CORS middleware, router registration", "High"],
        ["2", "backend/app/config.py", "Pydantic configuration, path management, .env loader", "High"],
        ["3", "backend/app/database.py", "SQLAlchemy database connection setup, engine, get_db dependency", "High"],
        ["4", "backend/app/dependencies/auth.py", "HTTPBearer JWT authentication dependency", "High"],
        ["5", "backend/app/models/user.py", "User ORM model", "High"],
        ["6", "backend/app/models/resume.py", "Resume ORM model with JSON skills and project columns", "High"],
        ["7", "backend/app/models/job_description.py", "JobDescription ORM model with JSON skills column", "Medium"],
        ["8", "backend/app/models/interview.py", "InterviewSession ORM model with adaptive state JSON", "High"],
        ["9", "backend/app/models/question.py", "InterviewQuestion, Answer, and AnswerEvaluation ORM models", "High"],
        ["10", "backend/app/services/interview_service.py", "Core interview orchestrator for creation, evaluation, adaptation", "High"],
        ["11", "backend/app/services/adaptive_engine.py", "Deterministic adaptive progression and skill selection engine", "High"],
        ["12", "backend/app/services/bloom.py", "Bloom's Revised Taxonomy definitions and helpers", "High"],
        ["13", "backend/app/services/evaluation_service.py", "Multi-factor answer scoring service (Groq + SBERT + Concept)", "High"],
        ["14", "backend/app/services/question_service.py", "Question generation and RAG ChromaDB singleton manager", "High"],
        ["15", "backend/app/services/jd_service.py", "Deterministic JD parser and priority skill mapping engine", "Medium"],
        ["16", "backend/app/services/resume_service.py", "Service wrapper for scripts/parse_resume.py", "Medium"],
        ["17", "backend/app/services/syllabus_rag_service.py", "Dynamic isolated RAG ingestion for academic course materials", "Medium"],
        ["18", "backend/app/services/syllabus_engine.py", "Syllabus topic selection with coverage constraints", "Medium"],
        ["19", "scripts/generate_question.py", "Skill taxonomy, domain mapping, ChromaDB retrieval, Groq generation", "High"],
        ["20", "scripts/parse_resume.py", "PyMuPDF resume text extractor and section parser", "High"],
        ["21", "scripts/ingest_vector_db.py", "Batch vector embedding and ChromaDB upsert pipeline", "High"],
        ["22", "scripts/chunk_text.py", "RecursiveCharacterTextSplitter and token filtering script", "Medium"],
        ["23", "frontend/src/App.jsx", "Frontend routing definitions and AppLayout wrapper", "High"],
        ["24", "frontend/src/context/AuthContext.jsx", "Global authentication context and token persistence", "High"],
        ["25", "frontend/src/services/api.js", "Axios client with Bearer token interceptor and 401 handler", "High"],
        ["26", "frontend/src/pages/Interview.jsx", "Real-time adaptive mock interview questioning page", "High"],
        ["27", "frontend/src/pages/Results.jsx", "Rich post-interview performance and recommendation page", "High"],
        ["28", "frontend/src/pages/ResumeUpload.jsx", "Resume and Job Description upload and skill mapping view", "High"],
        ["29", "frontend/src/pages/PerformanceDashboard.jsx", "Longitudinal multi-interview performance dashboard", "Medium"],
        ["30", "data/knowledge_stats.json", "Exact audit metadata for 402 chunks, 50 concepts, 8 domains", "High"]
    ]
    add_table_data(file_inventory[0], file_inventory[1:], [0.5, 2.5, 3.2, 1.0])

    add_heading_styled("29. Project Mentor Explanation", 1)
    add_heading_styled("2-Minute High-Level Pitch", 2)
    add_p("SmartInterview is an AI-powered technical mock interview platform that moves beyond static coding questionnaires. It extracts verified technical skills and personal projects from a candidate's resume, cross-references them with Job Description requirements, and dynamically conducts an interactive technical interview. What makes it unique is that an adaptive Python engine deterministically modulates question depth across Bloom's Taxonomy based on candidate answers, while a local vector database grounds questions in computer science literature to eliminate hallucinations. Candidates receive instant, multi-dimensional feedback and actionable preparation metrics.")

    add_heading_styled("5-Minute Architectural Overview", 2)
    add_p("Architecturally, SmartInterview connects a React 18 / Tailwind CSS frontend to a FastAPI backend backed by MySQL 8 and ChromaDB. When a candidate uploads their resume, PyMuPDF extracts their skills using a phrase-length descending canonical taxonomy. In the interview loop, questions are generated one at a time. The system maps the candidate's skill to its computer science domain, retrieves top-3 knowledge chunks from ChromaDB using SentenceTransformers, and prompts Groq's LLM to formulate a targeted question. When the candidate submits an answer, our hybrid evaluation engine scores technical accuracy, completeness, relevance, semantic similarity, and concept coverage. The resulting score updates the candidate's LearnerProfile, adjusting Bloom's cognitive complexity and selecting the next skill using a weakness-biased round-robin algorithm.")

    add_heading_styled("10-Minute Deep Technical Masterclass", 2)
    add_p("From an engineering perspective, SmartInterview enforces strict separation of concerns. Question generation and answer critique are delegated to high-speed LLM inference (Groq), but interview flow, skill selection, and difficulty adjustments are governed by deterministic Python state machines. Bloom's Taxonomy is modeled as an immutable 6-stage continuum. If a candidate scores >= 80, Bloom advances; if they score < 50, Bloom regresses. The database schema stores the complete serialized state machine inside a JSON column in interview_sessions, ensuring stateless backend scalability. For academic use cases, Syllabus Mode creates a temporary ChromaDB collection to isolate course materials, merging learned concepts into the permanent knowledge base only upon interview completion.")

    add_heading_styled("30. Limitations and Technical Debt", 1)
    add_bullet("Groq API rate limits (TPM/RPM) can throttle rapid back-to-back testing. Mitigated by exponential retry backoff.", "1. Single LLM Provider Dependency:")
    add_bullet("parse_resume.py relies on regex patterns and PyMuPDF text streams. Complex non-standard two-column PDFs can occasionally misclassify section boundaries.", "2. Resume Parsing Heuristics:")
    add_bullet("Storing JWT in browser localStorage is susceptible to Cross-Site Scripting (XSS). HttpOnly cookies with CSRF tokens would provide superior defense.", "3. LocalStorage JWT Storage:")
    add_bullet("No tab-scoped authentication isolation exists; logging in on a second tab overwrites tokens for the first tab.", "4. Multi-Tab Session Sharing:")
    add_bullet("Voice input/output, facial expression analysis, and real-time countdown timers are currently planned or not implemented.", "5. Missing Voice/Timer Features:")

    add_heading_styled("31. Final Master Workflow", 1)
    master_diagram = (
        "+-----------------------------------------------------------------------------------------+\n"
        "|                                     USER WORKFLOW                                       |\n"
        "|  1. Login / Register  ->  2. Resume & JD Upload  ->  3. Select Skills  ->  4. Start     |\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "                                             |\n"
        "                                             v\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "|                               ADAPTIVE QUESTION LOOP (Step n)                           |\n"
        "|  A. Select Next Skill (Weakness-Biased Round-Robin: min attempts -> lowest score)       |\n"
        "|  B. Map Skill -> Domain -> Query Template -> Vector Search (ChromaDB top-3 chunks)     |\n"
        "|  C. Construct Prompt (Bloom level + RAG context + Candidate Project + Dedup)            |\n"
        "|  D. Groq API Generation -> Question Displayed on React Interface                        |\n"
        "|  E. Candidate Writes & Submits Answer -> POST /questions/{qid}/answer                   |\n"
        "|  F. Hybrid Evaluation: Groq Critique + SBERT Cosine Similarity + Concept Coverage Math  |\n"
        "|  G. Overall Score Calculated (0-100) -> Stored in MySQL answer_evaluations              |\n"
        "|  H. Adaptive Decision:                                                                  |\n"
        "|       - Score >= 80 -> Advance Bloom (or Advance Difficulty if at Create)               |\n"
        "|       - Score 50-79 -> Maintain Level                                                   |\n"
        "|       - Score < 50  -> Regress Bloom (or Lower Difficulty if at Remember)               |\n"
        "|  I. Safeguard Check: Auto-stop if 3 consecutive fails < 40% at easy level               |\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "                                             |\n"
        "                                             v (Candidate clicks Finish)\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "|                                    RESULTS & ANALYTICS                                  |\n"
        "|  - Calculate Overall Average Score & Skill Performance Breakdowns                       |\n"
        "|  - Generate Personalized AI Recommendations                                             |\n"
        "|  - Display Interactive Milestone Visualizations on Results & Performance Pages          |\n"
        "+-----------------------------------------------------------------------------------------+"
    )
    add_code_block(master_diagram)

    add_heading_styled("32. Source of Truth & Evidence Map", 1)
    evidence_table = [
        ["Subsystem", "Primary Code Files", "Key Functions / Classes", "Verifiable Evidence"],
        ["Authentication", "routers/auth.py, services/auth_service.py", "register_user(), authenticate_user(), create_access_token()", "bcrypt hashing, JWT HS256 issuance, MySQL users table"],
        ["Resume Parsing", "scripts/parse_resume.py, services/resume_service.py", "parse_resume(), extract_candidate_name(), match_skills_in_text()", "PyMuPDF text extraction, generic header/footer cleaning"],
        ["Job Description", "routers/job_descriptions.py, services/jd_service.py", "parse_jd_file(), map_skills()", "Filters non-technical text, Group A & B skill mapping"],
        ["Vector Database", "scripts/ingest_vector_db.py, data/knowledge_stats.json", "ingest(), PersistentClient, collection.upsert()", "402 chunks, 50 concepts, 8 domains in ChromaDB technical_kb"],
        ["Question AI", "scripts/generate_question.py, question_service.py", "retrieve_rag_context(), build_prompt(), call_groq()", "SentenceTransformer embeddings, Groq API retry backoff"],
        ["Adaptive Engine", "services/adaptive_engine.py, services/bloom.py", "decide_next(), select_skill(), get_bloom_level()", "Deterministic score rules (80/50), weakness round-robin"],
        ["Answer Evaluation", "services/evaluation_service.py", "evaluate_answer(), _compute_semantic_similarity()", "30/20/20/15/15 weighted formula, SBERT cosine scaling"],
        ["Course Syllabus RAG", "services/syllabus_rag_service.py, syllabus_engine.py", "process_and_create_syllabus_rag(), select_next_topic()", "Isolated temporary Chroma collection, multi-topic coverage"],
        ["Database Schema", "setup_database.sql, setup_database_v2/3/4/5.sql", "CREATE TABLE users, resumes, sessions, questions, evaluations", "MySQL 8 tables, foreign keys, ON DELETE CASCADE, JSON types"],
        ["Frontend UI", "frontend/src/pages/Interview.jsx, Results.jsx", "handleSubmitAnswer(), handleFinishInterview(), loadResults()", "React 18 SPA, Axios interceptors, Tailwind CSS v4 styling"]
    ]
    add_table_data(evidence_table[0], evidence_table[1:], [1.2, 2.0, 2.0, 2.0])

    # ============================================================
    # EXECUTIVE SUMMARY
    # ============================================================
    add_heading_styled("Executive Summary", 1)
    add_p("The following 20 core points summarize the complete end-to-end operation of the SmartInterview platform:")
    
    exec_summary_points = [
        ("Platform Purpose: ", "SmartInterview is an automated AI-powered technical mock interview platform providing adaptive, resume-tailored, job-description-aligned, and syllabus-grounded interview practice."),
        ("Frontend Architecture: ", "Built as a single-page React 18 application with Vite 8, React Router v7, Tailwind CSS v4, and Lucide icons, reverse-proxied to the backend via Vite server configuration."),
        ("Backend Architecture: ", "Powered by FastAPI 0.115+ running on Uvicorn, structured into modular domain routers (/auth, /resumes, /job-descriptions, /interviews, /users)."),
        ("Database Persistence: ", "Utilizes MySQL 8.0 with SQLAlchemy 2.0 ORM, maintaining 7 relational tables with cascade deletes and JSON columns for flexible state storage."),
        ("Stateless Authentication: ", "Implements JWT bearer authentication with bcrypt password hashing; Axios automatically injects tokens and intercepts 401 Unauthorized responses."),
        ("Resume Ingestion: ", "PyMuPDF extracts candidate text while deterministic cleanup algorithms eliminate pagination, headers, and footers without LLM costs."),
        ("Canonical Taxonomy: ", "A phrase-length descending matching engine normalizes candidate skills into canonical names while avoiding partial substring errors."),
        ("Job Description Targeting: ", "Cross-matches resume skills against target job description skills, prioritizing overlap skills (Group A) and role gap skills (Group B)."),
        ("Production Knowledge Base: ", "ChromaDB technical_kb collection contains exactly 402 validated chunks across 50 concepts in 8 computer science domains."),
        ("Domain-Filtered RAG: ", "Resolves skills to technical domains and queries ChromaDB using SentenceTransformer all-MiniLM-L6-v2 (384 dimensions) to retrieve top-3 reference chunks."),
        ("Dynamic Question Generation: ", "Prompts Groq Cloud LLM (openai/gpt-oss-120b) with retrieved RAG context, Bloom cognitive level, and candidate project details, using exponential backoff retries."),
        ("One-at-a-Time Adaptive Loop: ", "Questions are generated dynamically one question at a time rather than in static upfront batches."),
        ("Multi-Factor Answer Evaluation: ", "Evaluates answers across 5 dimensions: Technical Correctness (30%), Completeness (20%), Relevance (20%), SBERT Cosine Similarity (15%), and Concept Coverage (15%)."),
        ("Bloom's Taxonomy Progression: ", "Models 6 cognitive levels (Remember, Understand, Apply, Analyze, Evaluate, Create) where scores >= 80 advance level and scores < 50 regress level."),
        ("Weakness-Biased Skill Selection: ", "Selects the next skill using an algorithm that prioritizes least-attempted skills, breaking ties by lowest average score."),
        ("Strict Auto-Stop Safeguard: ", "Automatically terminates an interview early only if 5+ questions are answered, difficulty is 'easy', and the candidate fails 3 consecutive questions (< 40%)."),
        ("Academic Syllabus Mode: ", "Allows uploading multi-file course materials (PDF, TXT, DOCX), creating an isolated temporary vector space and topic coverage progression engine."),
        ("Post-Interview Lifecycle Cleanup: ", "Syllabus mode merges learned course chunks into the permanent knowledge base and cleans up temporary collections upon interview completion."),
        ("Comprehensive Analytics: ", "Results and Performance Dashboard pages provide visual skill progress bars, cognitive milestones, and personalized AI recommendations."),
        ("Verified Implementation Truth: ", "Core mock interview, adaptive engine, vector RAG, and evaluation systems are fully implemented in code; voice STT/TTS and webcam facial analysis are planned future features.")
    ]

    for prefix, body in exec_summary_points:
        add_bullet(body, prefix)

    output_path = "SmartInterview_Master_Technical_Documentation.docx"
    doc.save(output_path)
    print(f"Master documentation successfully generated and saved to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_master_document()
