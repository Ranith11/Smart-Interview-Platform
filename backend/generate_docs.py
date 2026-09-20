from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\\Users\\user\\Desktop\\Smart-Interview-main\\SmartInterview_Implementation_Documentation.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=95, bottom=90, end=95):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, name="Aptos", size=10.2, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.2)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08

for name, size, before, after in (("Title", 22, 0, 12), ("Heading 1", 15, 15, 7), ("Heading 2", 12, 10, 5), ("Heading 3", 10.8, 8, 4)):
    style = styles[name]
    style.font.name = "Aptos Display" if name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Word's built-in Title style can carry a colored bottom border. This report
# intentionally uses title typography and whitespace only.
title_ppr = styles["Title"]._element.pPr
if title_ppr is not None:
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

footer = section.footer
add_page_number(footer.paragraphs[0])


def title(text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, "Aptos Display", 22, True, "000000")
    return p


def h1(text):
    return doc.add_paragraph(text, style="Heading 1")


def h2(text):
    return doc.add_paragraph(text, style="Heading 2")


def h3(text):
    return doc.add_paragraph(text, style="Heading 3")


def p(text="", bold_lead=None):
    para = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        r = para.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = para.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = para.add_run(text)
        set_run_font(r)
    return para


def bullet(text, level=0):
    para = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    para.paragraph_format.space_after = Pt(2)
    r = para.add_run(text)
    set_run_font(r)
    return para


def numbered(text):
    para = doc.add_paragraph(style="List Number")
    para.paragraph_format.space_after = Pt(2)
    r = para.add_run(text)
    set_run_font(r)
    return para


def code_block(text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.18)
    para.paragraph_format.right_indent = Inches(0.18)
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(6)
    for idx, line in enumerate(text.splitlines()):
        r = para.add_run(line)
        set_run_font(r, "Consolas", 8.4)
        if idx < len(text.splitlines()) - 1:
            r.add_break()
    return para


def table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    header = tbl.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.text = ""
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(str(text))
        set_run_font(run, size=8.4, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for row in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(tbl.rows) % 2 == 1:
                set_cell_shading(cell, "F4F8FB")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(str(value))
            set_run_font(run, size=8.25)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return tbl


def page_break():
    doc.add_page_break()


title("SmartInterview Implementation Documentation")
p("Technical documentation for the current SmartInterview main branch implementation.")
p("This document gives developers, reviewers, mentors, and testers a code-grounded view of the product, its runtime architecture, implemented workflows, data model, APIs, configuration, verification evidence, and current limitations.")
table(
    ["Item", "Value"],
    [
        ("Project", "SmartInterview"),
        ("Repository", "https://github.com/Ranith11/Smart-Interview-Platform.git"),
        ("Inspected branch", "main"),
        ("Inspected revision", "e4b5c74 - updated chromadb"),
        ("Documentation scope", "Current source code, configuration, data assets, and verification scripts"),
        ("Status terminology", "Implemented, Partially implemented, Documented or planned but not verified"),
    ],
    [1.5, 5.6],
)

h1("Contents")
for item in [
    "1 Product overview and feature inventory",
    "2 System architecture and end to end flow",
    "3 Frontend architecture",
    "4 Backend architecture and database design",
    "5 Authentication and resume intelligence",
    "6 Knowledge base RAG and domain normalization",
    "7 Normal Mode adaptive engine and Bloom taxonomy",
    "8 Question generation and answer evaluation",
    "9 Voice speech to text",
    "10 Syllabus Mode",
    "11 Interview lifecycle results history and dashboard",
    "12 API reference",
    "13 Project map configuration and running the project",
    "14 Error handling testing implementation status and limitations",
    "15 Extension points quick reference and final mental model",
]:
    bullet(item)

page_break()

h1("1 Product Overview and Feature Inventory")
h2("What SmartInterview is")
p("SmartInterview is a local full-stack technical mock interview application. In Normal Mode, a registered user uploads a resume, chooses parsed skills, receives AI-generated questions grounded by a technical knowledge base, submits text or spoken answers, and receives adaptive follow-up questions and evaluation feedback. In Syllabus Mode, the user uploads learning material, receives an interview on inferred syllabus topics, and can merge non-duplicate material into the permanent knowledge base when the interview completes.")
h2("Problem and users")
p("The product targets candidates practicing technical interviews and students practicing from supplied course material. It solves the problem of fixed, non-personalized questions by using resume-derived skills or uploaded material as the interview source, while persisting answers and evaluation results for later review.")
h2("Feature inventory")
table(
    ["Feature", "Purpose and processing", "Status", "Primary implementation"],
    [
        ("Authentication", "Register, bcrypt-hash passwords, login, issue and validate JWTs.", "IMPLEMENTED", "auth router, auth_service, AuthContext"),
        ("Resume intelligence", "Validate PDF, extract text, heuristically parse skills/projects/experience/education.", "IMPLEMENTED", "resumes router, parse_resume.py"),
        ("Normal Mode", "Resume-skill interview with one question generated at a time.", "IMPLEMENTED", "interview_service, InterviewSetup"),
        ("Permanent RAG", "Embed/query persisted technical_kb Chroma collection.", "IMPLEMENTED", "question_service, generate_question.py"),
        ("Adaptive engine", "Choose skill, Bloom level, difficulty, and type from latest score.", "IMPLEMENTED", "adaptive_engine.py"),
        ("Answer evaluation", "LLM scoring plus SBERT similarity and concept coverage.", "IMPLEMENTED", "evaluation_service.py"),
        ("Voice input", "Browser recording and backend faster-whisper transcription.", "IMPLEMENTED", "useSpeechToText, speech router/service"),
        ("Syllabus Mode", "Upload material, temporary RAG, sequential topic questions, merge/cleanup.", "PARTIALLY IMPLEMENTED", "syllabus services, interviews router"),
        ("History and results", "Persisted session question/answer/evaluation review.", "IMPLEMENTED", "interview_service, History, Results"),
        ("Performance dashboard", "Aggregated completed Normal Mode evaluations only.", "IMPLEMENTED", "get_user_performance, PerformanceDashboard"),
    ],
    [1.2, 3.25, 1.15, 1.55],
)
h2("Current implementation boundary")
bullet("Implemented behavior is determined by the repository source. Older week documents are historical material, not the source of truth where they conflict with current code.")
bullet("Normal Mode is user-open-ended in the UI but is automatically completed after 30 answered questions by backend safety logic.")
bullet("The UI does not expose API-supported Normal Mode difficulty, question type, or count configuration.")
bullet("Syllabus Mode still requires a valid user-owned resume because the backend validates resume ownership before branching by mode.")

h1("2 System Architecture and End to End Flow")
h2("High level architecture")
code_block("User browser\n  | React and Vite, sessionStorage JWT, Axios\n  v\nFastAPI API layer on port 8000\n  | routers, dependencies, services\n  +--> MySQL through SQLAlchemy\n  +--> ChromaDB technical_kb and temp_syllabus collections\n  +--> SentenceTransformer all-MiniLM-L6-v2\n  +--> Groq for generation, topic extraction, and evaluation\n  +--> faster-whisper for speech transcription")
h2("Complete end to end flow")
code_block("Register or login\n  -> JWT stored in sessionStorage\n  -> protected React route\n  -> upload resume or syllabus material\n  -> configure mode and start session\n  -> backend persists session and generates first question\n  -> user provides text or voice answer\n  -> answer is persisted and evaluated\n  -> deterministic progression chooses next state\n  -> next question or completion\n  -> results, history, and performance views")
h2("Execution chains")
table(
    ["Workflow", "Actual code path"],
    [
        ("Register and login", "Register/Login page -> AuthContext -> Axios -> auth router -> auth_service -> User table -> JWT -> sessionStorage."),
        ("Resume upload", "ResumeUpload -> /resumes/upload -> file validation -> unique storage file -> resume_service -> parse_resume.py -> Resume row."),
        ("Normal question", "InterviewSetup -> /interviews/start -> create_interview_session -> AdaptiveState -> generate_single_question -> RAG retrieval -> Groq -> InterviewQuestion."),
        ("Answer loop", "Interview -> answer endpoint -> Answer row -> evaluate_answer -> AnswerEvaluation -> decide_next -> next question."),
        ("Syllabus upload", "InterviewSetup -> /interviews/upload-syllabus -> extraction/chunking -> temporary Chroma collection -> subject/topic LLM responses."),
        ("Speech", "Speech hook -> multipart /speech/transcribe -> temp webm -> faster-whisper -> final transcript appended to answer field."),
    ],
    [1.55, 5.6],
)

h1("3 Frontend Architecture")
h2("Frontend technology and entry points")
table(
    ["Area", "Implementation"],
    [
        ("Framework", "React mounted in Strict Mode by src/main.jsx"),
        ("Build", "Vite with React and Tailwind plugins"),
        ("Development port", "5173"),
        ("API proxy", "/api proxies to http://localhost:8000"),
        ("Routing", "react-router-dom BrowserRouter"),
        ("State", "React local state and AuthContext only; no Redux, Zustand, or React Query"),
        ("Styling", "Tailwind CSS v4 plus reusable btn, card, input, label, and badge classes in index.css"),
    ],
    [1.5, 5.65],
)
h2("Routes")
table(
    ["Route", "Protection", "Component and purpose"],
    [
        ("/", "Public", "Landing.jsx - product entry page"),
        ("/login and /register", "Public", "Login.jsx and Register.jsx"),
        ("/dashboard", "Protected", "Dashboard.jsx - activity snapshot"),
        ("/resume", "Protected", "ResumeUpload.jsx - active resume management"),
        ("/setup", "Protected", "InterviewSetup.jsx - Normal Mode"),
        ("/setup?mode=syllabus", "Protected", "InterviewSetup.jsx - Syllabus Mode"),
        ("/interview/:id", "Protected", "Interview.jsx - active interview"),
        ("/results/:id", "Protected", "Results.jsx - detailed result review"),
        ("/history, /profile, /performance", "Protected", "History, Profile, and PerformanceDashboard"),
    ],
    [1.55, 1.25, 4.35],
)
h2("Authentication state and API client")
p("AuthContext is the shared client-side state boundary. It removes legacy localStorage keys, restores sessionStorage credentials, verifies the token through GET /auth/me, and exposes user, loading, login, register, logout, and isAuthenticated. The Axios client attaches a Bearer token from sessionStorage to every request and clears session state after a 401 response.")
h2("Important frontend file map")
table(
    ["File", "Responsibility", "Used by"],
    [
        ("src/App.jsx", "Routes, AuthProvider, protected layout and Sidebar.", "Application root"),
        ("src/context/AuthContext.jsx", "Session persistence and auth operations.", "All auth-aware pages/components"),
        ("src/services/api.js", "Axios base URL, token injection, 401 redirect.", "All API-using pages/hooks"),
        ("src/pages/InterviewSetup.jsx", "Skills, syllabus staging/upload, start payload.", "Routes /setup"),
        ("src/pages/Interview.jsx", "Question, answer, evaluation delay, finish action, speech hook.", "Route /interview/:id"),
        ("src/pages/Results.jsx", "Session result and score breakdown rendering.", "Route /results/:id"),
        ("src/pages/PerformanceDashboard.jsx", "Normal Mode analytics charts and summaries.", "Route /performance"),
        ("src/hooks/useSpeechToText.js", "Recorder, VAD, request versioning, transcript callback.", "Interview.jsx"),
        ("src/components/ProtectedRoute.jsx", "Client navigation guard and loading spinner.", "Protected routes"),
        ("src/components/Sidebar.jsx", "Authenticated desktop/mobile navigation.", "AppLayout"),
    ],
    [2.0, 3.45, 1.7],
)
h2("Frontend state transitions")
code_block("Auth loading -> authenticated or redirected to login\nInterview loading -> current unanswered question -> answer draft\nAnswer submit -> submitting -> evaluation displayed -> next question after 3 seconds\nCompletion -> results route\nSpeech inactive -> recording -> transcribing -> inactive")

h1("4 Backend Architecture and Database Design")
h2("FastAPI application")
p("backend/app/main.py creates SmartInterview API version 7.0.0, registers auth, resumes, interviews, users, and speech routers, and exposes GET /api/health. CORS allows only Vite localhost origins.")
h2("Backend layer diagram")
code_block("router\n  -> get_current_user and get_db dependencies\n  -> service orchestration\n  -> SQLAlchemy models / ChromaDB / Groq / SBERT / faster-whisper\n  -> Pydantic or dictionary response")
h2("Important backend files")
table(
    ["File", "Key responsibilities", "Called by"],
    [
        ("app/config.py", "Loads environment values and establishes paths.", "All infrastructure services"),
        ("app/database.py", "Engine, SessionLocal, declarative Base, get_db lifecycle.", "Routers and models"),
        ("app/dependencies/auth.py", "HTTPBearer JWT decoding and user lookup.", "Protected router endpoints"),
        ("app/routers/interviews.py", "Interview, syllabus, history, result, answer, completion endpoints.", "Frontend API client"),
        ("app/services/interview_service.py", "Session creation, answer loop, aggregation.", "Interview router"),
        ("app/services/question_service.py", "Lazy singleton setup and normal/syllabus question generation.", "Interview service"),
        ("app/services/evaluation_service.py", "LLM evaluation, similarity, final score.", "Interview service"),
        ("app/services/adaptive_engine.py", "Deterministic Normal Mode decision engine.", "Interview service"),
        ("app/services/syllabus_rag_service.py", "Syllabus extraction, temp collections, merge and delete.", "Interview router/service"),
    ],
    [2.05, 3.55, 1.55],
)
h2("Entity relationship design")
code_block("User 1 -> many Resume\nUser 1 -> many InterviewSession\nResume 1 -> many InterviewSession\nInterviewSession 1 -> many InterviewQuestion\nInterviewSession 1 -> many Answer\nInterviewQuestion 1 -> 0..1 Answer\nAnswer 1 -> 0..1 AnswerEvaluation\nInterviewQuestion 1 -> 0..1 AnswerEvaluation at ORM level")
h2("Database entities")
table(
    ["Model table", "Important fields", "Relationships and purpose"],
    [
        ("users", "id, name, unique email, password_hash, timestamps", "Owns resumes and sessions; identity record."),
        ("resumes", "user_id, filename/path, skills/projects/experience/education JSON, raw_text", "Candidate profile; sessions reference it."),
        ("interview_sessions", "difficulty, type/count, status, adaptive/syllabus state, mode, completion data", "Lifecycle record for one interview."),
        ("interview_questions", "session, ordinal, skill/topic, text, RAG/project metadata, Bloom", "Generated questions for a session."),
        ("answers", "unique question_id, session/user IDs, answer_text", "One persisted user answer per question."),
        ("answer_evaluations", "five component scores, overall, feedback and concept arrays", "One evaluation per answer through unique answer_id."),
    ],
    [1.55, 3.0, 2.6],
)
h2("Constraints and stored JSON")
bullet("users.email, answers.question_id, and answer_evaluations.answer_id are unique. Foreign keys use ON DELETE CASCADE.")
bullet("JSON state includes resume profile fields, selected skills, AdaptiveState, SyllabusState, final recommendations, RAG metadata, project context, and qualitative evaluation arrays.")
bullet("The base schema is setup_database.sql. v2 adds Bloom/adaptive/evaluation support, v3 adds completion_reason, and migrate_db.py adds mode/syllabus columns with a hard-coded local connection string.")

h1("5 Authentication and Resume Intelligence")
h2("Authentication and authorization")
code_block("Register request -> compare confirmation -> bcrypt hash -> User row -> JWT\nLogin request -> bcrypt.checkpw -> JWT\nBearer request -> jose JWT decode -> subject user ID -> current User query")
table(
    ["Control", "Verified implementation"],
    [
        ("Password handling", "bcrypt hashing, 72-byte input cap before hash or verify."),
        ("JWT", "HS256 by default; sub user ID and expiration, default 1440 minutes."),
        ("Authorization", "HTTPBearer dependency resolves the current user for protected endpoints."),
        ("Ownership", "Resume/session queries include authenticated user ID."),
        ("Client token storage", "sessionStorage; tab-scoped; global Axios 401 clearing."),
        ("Not implemented", "Refresh tokens, revocation, roles, rate limiting, account recovery, email verification."),
    ],
    [1.8, 5.35],
)
h2("Resume intelligence pipeline")
code_block("PDF -> PyMuPDF text extraction -> whitespace normalization\n-> section detection -> skill/project/experience/education heuristics\n-> fallback keyword scan for skills -> Resume JSON fields")
p("The parser uses heading patterns for Skills, Projects, Experience, and Education. Skills are split using common separators; projects and experience are inferred from line heuristics; project technologies come from a fixed technology-keyword list. The implementation is heuristic, not OCR or LLM resume parsing.")
table(
    ["Validation or behavior", "Actual result"],
    [
        ("File type", "Resume endpoint accepts PDF filename extension only."),
        ("Size", "Maximum 5 MB; empty files rejected."),
        ("Storage", "Generated user-ID plus UUID filename, not original name."),
        ("No skills", "Upload rejected with 422."),
        ("Replacement", "Existing user resumes are deleted; related sessions cascade-delete."),
        ("Boundary", "Image-only PDFs have no OCR path. Parser can use sys.exit for no text."),
    ],
    [1.75, 5.4],
)

h1("6 Knowledge Base RAG and Domain Normalization")
h2("Permanent knowledge base")
table(
    ["Metric", "Verified value"],
    [
        ("Processed JSON files", "50"),
        ("Canonical domains", "8: cn, dbms, design-patterns, dsa, ml-dl, oop, os, system-design"),
        ("Concepts", "50"),
        ("Unique chunks", "402"),
        ("Token range", "50 to 385"),
        ("Embedding model", "all-MiniLM-L6-v2"),
        ("Persistent collection", "technical_kb in chroma_db"),
    ],
    [2.0, 5.15],
)
h2("RAG ingestion and retrieval")
code_block("raw HTML -> extract_text.py -> clean_text.py -> chunk_text.py\n-> data/processed/domain/concept_chunks.json -> ingest_vector_db.py\n-> Chroma technical_kb\n\nskill -> resolve_skill_domains -> build_rag_query -> SBERT embedding\n-> Chroma domain filter when mapped -> top 3 chunks -> Groq prompt")
p("chunk_text.py uses RecursiveCharacterTextSplitter with 1024-character chunks and 200-character overlap. It retains chunks only from 50 to 400 cl100k_base tokens. Ingestion validates required metadata and upserts batches of 100. Retrieval returns three nearest chunks; there is no current relevance distance cutoff.")
h2("Domain normalization")
table(
    ["Canonical domain", "Examples normalized into it"],
    [
        ("dbms", "database, database management systems, SQL, DBMS"),
        ("oop", "object-oriented programming, OOP"),
        ("dsa", "data structures, algorithms, DSA"),
        ("design-patterns", "design pattern or design patterns"),
        ("system-design", "system design or systems design"),
        ("os", "operating system or operating systems"),
        ("cn", "networking, computer networks, CN"),
        ("ml-dl", "machine learning, deep learning, artificial intelligence, ML, AI"),
    ],
    [2.0, 5.15],
)
p("SKILL_DOMAIN_MAP in scripts/generate_question.py maps skills such as Java, Python, SQL, React, Docker, cloud platforms, TensorFlow, and Git to one or more canonical domains. Exact match is tried first, then a partial match. An unmapped skill uses unfiltered retrieval.")
h2("RAG validation evidence")
p("The read-only validate_repository.py run confirmed 402 unique chunks over 50 files and eight domains. It produced five warnings for residual web-artifact phrases: three subscribe warnings, one related articles warning, and one cookie warning. The evaluator script and historical metric claims exist but were not run for this document.")

h1("7 Normal Mode Adaptive Engine and Bloom Taxonomy")
h2("Normal Mode workflow")
code_block("validate owned resume -> select valid resume skills -> initialize AdaptiveState\n-> first decision at Remember -> create session -> generate first question\n-> submit answer -> evaluate -> persist evaluation\n-> deterministic next decision -> generate next question\n-> manual finish or safety-limit completion")
h2("Adaptive state and decision rules")
table(
    ["State or rule", "Actual behavior"],
    [
        ("Initial state", "First selected skill; Bloom remember; configured/default medium difficulty."),
        ("Score >= 80", "Advance Bloom; if Create already reached, raise difficulty."),
        ("Score 50 to 79", "Maintain Bloom and difficulty."),
        ("Score < 50", "Regress Bloom; if Remember, lower difficulty."),
        ("Skill selection", "Fewest attempts, then lowest average score among ties."),
        ("Question type", "First preferred type mapped from resulting Bloom level."),
        ("Deduplication", "Previous successful question text is included in future prompts."),
        ("Completion", "Manual finish or questions_answered >= MAX_ADAPTIVE_QUESTIONS, which is 30."),
    ],
    [2.0, 5.15],
)
h2("Bloom progression")
code_block("Remember -> Understand -> Apply -> Analyze -> Evaluate -> Create\n  recall      explain     use       compare    judge       design\n\nScore below 50 moves one level backward. Score at least 80 moves one level forward.")
table(
    ["Bloom level", "Selected question type", "Prompt intent"],
    [
        ("Remember", "conceptual", "Recall or define a concept."),
        ("Understand", "conceptual", "Explain or interpret a concept."),
        ("Apply", "practical", "Use knowledge to solve a problem."),
        ("Analyze", "technical_reasoning", "Compare or break down approaches."),
        ("Evaluate", "scenario", "Judge or justify technical decisions."),
        ("Create", "project", "Design or propose a solution."),
    ],
    [1.4, 1.8, 3.95],
)
p("Bloom is applied by Normal Mode question generation. Syllabus questions do not receive Bloom fields or Bloom prompt augmentation.")
h2("Recommendations")
table(
    ["Average score", "Recommendation priority and direction"],
    [
        ("Below 40", "High: improve foundational concepts before advanced topics."),
        ("40 to 59", "Medium: review core concepts and practical application."),
        ("60 to 79", "Low: practice more complex analysis and design."),
        ("80 or above", "Positive: continue advanced study, naming reached Bloom level."),
    ],
    [1.7, 5.45],
)

h1("8 Question Generation and Answer Evaluation")
h2("Question generation")
p("question_service.py lazily initializes the embedding model, technical_kb Chroma collection, Groq client, configured model, and prompt JSON. It delegates core RAG/prompt functions to scripts/generate_question.py.")
code_block("skill/topic + difficulty + type + optional project context + top 3 RAG chunks\n+ prior questions + optional Bloom guidance\n-> Groq question -> InterviewQuestion persistence")
table(
    ["Prompt safeguard", "Current prompt instruction"],
    [
        ("Question count", "Generate exactly one clear concise question."),
        ("Focus", "Ask about one primary technical concept and keep the selected skill central."),
        ("Candidate context", "Do not invent experience or technologies; project context is optional personalization."),
        ("Question shape", "Do not ask multi-part questions and do not provide the answer."),
        ("Difficulty", "Easy 10-25 words; medium 15-35; hard 20-45."),
        ("Retry behavior", "Up to three Groq attempts with 5, 10, and 20 second backoff; failed output becomes [GENERATION FAILED]."),
    ],
    [1.65, 5.5],
)
h2("Answer evaluation pipeline")
code_block("Question and RAG reference + candidate answer\n  -> Groq JSON evaluation\n  -> SBERT cosine similarity\n  -> expected/found concept coverage check\n  -> weighted overall score\n  -> AnswerEvaluation persistence")
table(
    ["Score", "Source", "Weight"],
    [
        ("Technical correctness", "Groq", "30 percent"),
        ("Completeness", "Groq", "20 percent"),
        ("Relevance", "Groq", "20 percent"),
        ("Semantic similarity", "SentenceTransformer cosine similarity", "15 percent"),
        ("Concept coverage", "Groq plus calculated expected/found ratio when available", "15 percent"),
    ],
    [2.0, 3.7, 1.45],
)
p("Overall score = 0.30 technical + 0.20 completeness + 0.20 relevance + 0.15 semantic similarity + 0.15 concept coverage. Scores are clamped to 0-100. Answers shorter than five characters return zero overall score without an LLM request.")
p("Similarity uses the answer embedding against question text plus up to 1,500 characters of RAG reference. Cosine similarity is scaled from the approximate 0.1-0.7 range to 0-100. It is one signal and is not treated as proof of technical correctness.")

h1("9 Voice Speech to Text")
h2("Implemented speech flow")
code_block("getUserMedia -> AudioContext and AnalyserNode -> VAD\n-> MediaRecorder chunks every second -> interim/final multipart request\n-> /api/speech/transcribe -> temporary webm -> faster-whisper\n-> final text appended to answer textarea")
table(
    ["Component", "Verified behavior"],
    [
        ("VAD", "Average frequency value above 10 means speech; 1,000 ms silence finalizes a segment."),
        ("Recorder", "MediaRecorder selects a supported webm/mp4/ogg MIME type and emits 1-second chunks."),
        ("Interim protection", "Backpressure avoids overlapping interim requests; versions and segment IDs discard stale responses."),
        ("Final fallback", "If final transcription fails, last interim text is appended when available."),
        ("Server", "Requires auth, rejects empty or over-5-MB input, writes a temporary webm, always tries cleanup."),
        ("Model", "Lazy faster-whisper WhisperModel with configurable model, device, and compute type."),
    ],
    [1.85, 5.3],
)
p("The endpoint does not actually reject unsupported MIME types because its content-type conditional only executes pass. Browser microphone behavior was not runtime-tested for this report.")

h1("10 Syllabus Mode")
h2("Syllabus workflow")
code_block("PDF/TXT/DOCX uploads -> validation and unique file storage\n-> extract text -> 1000-character chunks with 200 overlap\n-> subject inference -> LLM topic extraction\n-> temp_syllabus UUID collection -> sequential topic interview\n-> score recording -> completion -> permanent-KB merge -> temp collection deletion")
h2("Upload and temporary RAG")
table(
    ["Aspect", "Actual implementation"],
    [
        ("Allowed files", "PDF, TXT, DOCX; multiple uploads supported."),
        ("Limit", "20 MB per file; empty files rejected."),
        ("Extraction", "PyMuPDF for PDF, UTF-8 replacement decoding for TXT, python-docx for DOCX."),
        ("Subject", "Groq inference from first file sample, with filename fallback."),
        ("Topics", "Groq returns a cleaned unique list; source asks for 4-16 high-level topics."),
        ("Collection", "temp_syllabus_<uuid>; embeddings and metadata are stored in ChromaDB."),
        ("Metadata", "normalized domain, LLM-extracted concept, source filename, source label."),
    ],
    [1.65, 5.5],
)
h2("Syllabus state and completion")
table(
    ["State field", "Meaning"],
    [
        ("selected_topics", "Topics selected for this session."),
        ("questions_per_topic", "Fixed quota calculated from request count or default 10."),
        ("current_topic_index", "Position in deterministic topic order."),
        ("questions_answered_in_topic and total", "Progress counters."),
        ("previous_questions", "Prompt de-duplication history."),
        ("topic_scores", "Score lists used for final topic recommendations."),
    ],
    [2.35, 4.8],
)
h2("Temporary-to-permanent merge")
p("For every temporary chunk, merge_temporary_to_permanent queries the closest permanent vector, calculates cosine similarity, treats similarity greater than 0.85 as a duplicate, and upserts non-duplicates into technical_kb. Natural or manual completion then attempts to delete the temporary collection. Merge errors are logged rather than propagated to the user.")
h2("Normal Mode and Syllabus Mode comparison")
table(
    ["Aspect", "Normal Mode", "Syllabus Mode"],
    [
        ("Source", "Resume skills and permanent technical KB", "Uploaded reference material and temp KB"),
        ("State", "AdaptiveState", "SyllabusState"),
        ("Progression", "Score-adaptive Bloom/difficulty/skill", "Sequential topic quota; scores do not alter order"),
        ("Question context", "Permanent technical_kb, filtered where mapped", "temp_syllabus collection, top 3 topic-query chunks"),
        ("Completion", "Manual or 30-answer safety limit", "All topic quotas or manual completion"),
        ("Knowledge merge", "None", "Non-duplicate temp chunks can enter permanent KB"),
    ],
    [1.35, 2.8, 2.95],
)
h2("Syllabus Mode implementation boundaries")
bullet("A resume is still required by backend session creation even though syllabus setup does not show that requirement.")
bullet("The current frontend auto-selects all inferred topics. toggleTopic exists but no selection UI uses it.")
bullet("With null question_count from the UI, questions per topic is max(1, 10 divided by topic count) while stored session question_count is topic count times 2; displayed and actual totals can differ.")
bullet("Evaluation retrieves full RAG context only from permanent technical_kb, not the temporary syllabus collection. Syllabus evaluations usually lack uploaded-source reference text and semantic similarity is zero before merge.")
bullet("Uploads and temp collections without session completion have no explicit cleanup endpoint or expiry path.")

h1("11 Interview Lifecycle Results History and Dashboard")
h2("Lifecycle")
code_block("in_progress -> question generated -> answer persisted -> evaluation persisted\n-> next question or completed\n\nSession status values: in_progress, completed, abandoned\nCurrent services create in_progress/completed. No route was found that sets abandoned.")
table(
    ["Completion reason", "When assigned"],
    [
        ("manual", "User calls POST /api/interviews/{id}/complete."),
        ("assessment_complete", "All Syllabus Mode topic quotas are answered."),
        ("max_questions_safety_limit", "Normal Mode reaches 30 answered questions."),
    ],
    [2.25, 4.9],
)
h2("Results and feedback")
p("get_session_results builds question-by-question output from MySQL: question text, answer text, evaluation components, feedback, strengths, weaknesses, expected/found concepts, calculated per-skill averages, calculated overall average, stored recommendations, and persisted Bloom progression. Results.jsx renders the score summary, skill bars, recommendations, Bloom history when present, and expandable question review.")
h2("Interview history")
p("History returns every session owned by the user in newest-first order. It counts answers for each session. It calculates and returns average_score only for truthy is_adaptive sessions, so syllabus sessions can have evaluations without a displayed history average.")
h2("Performance dashboard")
table(
    ["Metric", "Backend calculation"],
    [
        ("Included sessions", "Completed current-user sessions where mode equals normal. Syllabus mode is excluded."),
        ("Overall average", "Mean of stored overall_score values over included evaluations."),
        ("Skill performance", "Average/count/highest Bloom order for each actually tested skill."),
        ("Strengths", "Skills with average at least 75."),
        ("Weak areas", "Skills with average below 50."),
        ("Bloom performance", "Average evaluation score grouped by Bloom label."),
        ("Recent interviews", "Completed normal sessions in descending completion order."),
        ("Recommendations", "Stored session recommendations deduplicated by skill, up to 10."),
    ],
    [2.05, 5.1],
)
p("Dashboard statistics from get_user_stats differ from the performance dashboard: they count all session modes and all user evaluations. The performance dashboard can return data for a completed Normal Mode session without evaluations, producing a zero average.")

h1("12 API Reference")
table(
    ["Method", "Endpoint", "Auth", "Purpose"],
    [
        ("GET", "/api/health", "No", "Health response with status and service name."),
        ("POST", "/api/auth/register", "No", "Create account and return access token/user."),
        ("POST", "/api/auth/login", "No", "Authenticate and return access token/user."),
        ("GET", "/api/auth/me", "Yes", "Return current user."),
        ("POST", "/api/resumes/upload", "Yes", "Upload and parse PDF resume."),
        ("GET", "/api/resumes/current", "Yes", "Get current user resume."),
        ("DELETE", "/api/resumes/{id}", "Yes", "Delete owned resume."),
        ("GET", "/api/users/profile", "Yes", "User, latest resume, and aggregate stats."),
        ("PUT", "/api/users/profile", "Yes", "Update optional name."),
        ("GET", "/api/users/stats", "Yes", "Aggregate counts and average score."),
        ("GET", "/api/users/performance", "Yes", "Completed Normal Mode performance aggregation."),
        ("POST", "/api/interviews/upload-syllabus", "Yes", "Create temporary syllabus RAG and return topic data."),
        ("POST", "/api/interviews/start", "Yes", "Create session and first question."),
        ("GET", "/api/interviews/history", "Yes", "List current user sessions."),
        ("GET", "/api/interviews/{id}", "Yes", "Session state with questions and evaluations."),
        ("GET", "/api/interviews/{id}/results", "Yes", "Rich result payload."),
        ("GET", "/api/interviews/{id}/questions/{number}", "Yes", "Legacy numbered-question read."),
        ("POST", "/api/interviews/{id}/questions/{question_id}/answer", "Yes", "Persist/evaluate answer and create next question."),
        ("POST", "/api/interviews/{id}/complete", "Yes", "Manually complete session."),
        ("POST", "/api/speech/transcribe", "Yes", "Transcribe multipart audio and return text."),
    ],
    [0.65, 3.2, 0.55, 2.75],
)
h2("Important request and response contracts")
code_block("StartInterviewRequest\n  resume_id: integer\n  difficulty: easy | medium | hard | null\n  question_type: string | null\n  question_count: integer | null\n  selected_skills: list[string] | null\n  mode: normal | syllabus\n  syllabus_id: string | null\n  selected_topics: list[string] | null\n\nAnswerRequest\n  answer_text: string, maximum 5000 characters")
p("The adaptive answer response returns evaluation, optional next_question, is_complete, questions_answered, questions_remaining, current_bloom_level, and current_difficulty. Standard error paths include Pydantic 422 validation, 401 authentication, 404 ownership/resource failures, 400 domain validation, and caught provider failures as 500.")

h1("13 Project Map Configuration and Running the Project")
h2("Practical repository map")
code_block("frontend/\n  src/pages, src/components, src/context, src/hooks, src/services\n  package.json, vite.config.js\n\nbackend/\n  app/models, app/schemas, app/routers, app/services, dependencies\n  setup_database.sql, setup_database_v2.sql, setup_database_v3.sql\n  e2e and verification scripts\n\nscripts/\n  content download, extraction, cleaning, chunking, ingestion, retrieval validation\n\ndata/\n  raw source captures, processed chunks, retrieval query set, knowledge statistics\n\nprompts/\n  question_generation.json, evaluation_prompts.json\n\nchroma_db/\n  persisted ChromaDB storage")
h2("Configuration")
table(
    ["Variable", "Purpose"],
    [
        ("DATABASE_URL", "SQLAlchemy MySQL connection URL."),
        ("JWT_SECRET_KEY", "JWT signing key."),
        ("JWT_ALGORITHM", "JWT algorithm; default HS256."),
        ("JWT_EXPIRATION_MINUTES", "JWT lifetime; default 1440."),
        ("GROQ_API_KEY", "Groq credential."),
        ("GROQ_MODEL", "Groq model; default openai/gpt-oss-120b."),
        ("STT_MODEL_SIZE", "faster-whisper model; default tiny.en."),
        ("STT_DEVICE", "Speech model device; default cpu."),
        ("STT_COMPUTE_TYPE", "Speech model computation type; default int8."),
    ],
    [2.15, 5.0],
)
h2("Run the project")
numbered("Create the MySQL smartinterview database using the baseline setup SQL, then apply the current migrations represented in setup_database_v2.sql, setup_database_v3.sql, and migrate_db.py as appropriate for the existing database.")
numbered("Create a root .env from .env.example and configure the database, JWT, Groq, and optional STT settings. Do not use placeholder or development secrets for deployment.")
numbered("Run start.bat. It creates backend/venv and installs backend requirements on first use, starts uvicorn app.main:app --reload from backend, installs frontend node modules on first use, and runs npm run dev.")
numbered("Verify GET http://localhost:8000/api/health and open http://localhost:5173.")
p("The backend requirements include FastAPI, Uvicorn, SQLAlchemy, PyMySQL, python-jose, bcrypt support, multipart upload support, SentenceTransformers, ChromaDB, PyMuPDF, python-docx, and faster-whisper. Frontend package.json defines dev, build, and preview scripts.")

h1("14 Error Handling Testing Implementation Status and Limitations")
h2("Error handling")
table(
    ["Area", "Verified behavior"],
    [
        ("Authentication", "Invalid, expired, malformed, or userless JWT yields 401."),
        ("Resume upload", "Rejects non-PDF, empty, oversized, unparseable, or no-skill resumes."),
        ("Syllabus upload", "Rejects empty/missing/unsupported/oversized files; extraction errors may be logged and skipped."),
        ("Interview ownership", "Session/resume lookup includes current user; missing records yield 404 or service ValueError."),
        ("Duplicate answers", "Adaptive submission rejects an already populated answer."),
        ("LLM evaluation", "Malformed/failed LLM output becomes safe default zero scores and feedback."),
        ("RAG/evaluation lookup", "Failures are caught for re-retrieval and evaluation can continue with no reference context."),
        ("Speech", "Temporary file cleanup runs in finally; transcription failures map to 500."),
        ("Frontend", "Per-page local spinners/errors; Axios globally handles 401; no error boundary."),
    ],
    [1.65, 5.5],
)
h2("Testing and verification inventory")
table(
    ["Test or check", "What it verifies", "Status"],
    [
        ("scripts/validate_repository.py", "Chunk metadata, ID uniqueness, bounds, content warnings.", "Executed read-only: 402 unique chunks; five warnings."),
        ("scripts/validate_chunks.py", "One chunk JSON quality report.", "Present; not run."),
        ("scripts/test_retrieval.py", "Manual Chroma top-k retrieval.", "Present; not run."),
        ("scripts/evaluate_retrieval.py", "Hit@1/3/5 and MRR over 33 queries.", "Present; not run."),
        ("backend/e2e_api_test.py", "Syllabus HTTP flow, merge, cleanup.", "Present; not run because it needs live services and writes data."),
        ("backend/test_optimization.py", "PDF/TXT/DOCX syllabus upload.", "Present; not run because it creates test data."),
        ("backend/test_perf_db.py", "Normal-only performance filtering.", "Present; not run because it may create mock rows."),
        ("Frontend test runner", "Unit/component/browser tests.", "No test configuration or scripts found."),
    ],
    [2.0, 3.45, 1.7],
)
h2("Current implementation status")
table(
    ["Feature", "Status", "Evidence and note"],
    [
        ("JWT authentication", "IMPLEMENTED", "Auth router/service/dependency and AuthContext."),
        ("Resume PDF parsing", "IMPLEMENTED", "Resume router and parser; heuristic only."),
        ("Normal RAG/adaptive interview", "IMPLEMENTED", "Question, evaluation, interview, adaptive services."),
        ("Bloom taxonomy", "IMPLEMENTED", "Normal Mode only."),
        ("Voice transcription", "IMPLEMENTED", "Hook, speech endpoint, faster-whisper; runtime browser verification not captured."),
        ("Syllabus temporary RAG", "IMPLEMENTED", "Upload, collection generation, topic flow, merge service."),
        ("Syllabus selection and grounding", "PARTIALLY IMPLEMENTED", "No visible topic customization; evaluation misses temp source context."),
        ("Dashboard", "IMPLEMENTED", "Completed Normal Mode only; syllabi excluded."),
        ("Automated frontend tests", "NOT VERIFIED", "No test configuration found."),
        ("Background job processing", "NOT IMPLEMENTED", "AI/STT calls are synchronous request work."),
    ],
    [2.0, 1.5, 3.65],
)
h2("Known limitations")
table(
    ["Limitation", "Impact", "Current behavior"],
    [
        ("Resume dependency in Syllabus Mode", "Syllabus-only start fails.", "resume_id ownership is verified before mode branch."),
        ("No Normal Mode configuration controls", "API settings are unavailable to UI users.", "UI submits null difficulty/type/count."),
        ("Hidden 30-question cap", "Open-ended interview is bounded.", "Normal Mode completes at 30 answers."),
        ("Syllabus source not used in evaluation", "Feedback lacks uploaded-material context.", "Evaluation looks in permanent technical_kb only."),
        ("No temp resource expiry", "Unused files/collections may accumulate.", "Cleanup occurs only on completed/manual interview."),
        ("Heuristic resume parsing", "Unusual formats can parse poorly.", "No OCR/LLM parser."),
        ("MIME check does not enforce", "Non-audio input may reach transcription.", "Unsupported type branch only passes."),
        ("No frontend tests", "UI regressions have limited evidence.", "No test runner configured."),
    ],
    [2.0, 2.45, 2.7],
)

h1("15 Extension Points Quick Reference and Final Mental Model")
h2("Logical extension points")
bullet("Allow Syllabus Mode without a resume and make resume_id nullable for those sessions.")
bullet("Render topic selection controls and Normal Mode configuration controls already represented by the API schema.")
bullet("Retrieve temporary syllabus RAG chunks during evaluation, then preserve source-aware feedback.")
bullet("Add an explicit discard endpoint, lifecycle TTL, and scheduled cleanup for abandoned syllabus artifacts.")
bullet("Move slow LLM, embedding, and speech operations into background jobs with status reporting and provider timeouts.")
bullet("Use managed database migrations and add the intended uniqueness constraint for question evaluations.")
bullet("Add API, frontend component, browser voice, and end-to-end automated tests.")
bullet("Add production security hardening: mandatory secrets, token refresh/revocation, rate limits, and observability.")
h2("Quick reference")
table(
    ["Area", "Reference"],
    [
        ("Frontend", "React, Vite, React Router, Axios, Tailwind CSS, Lucide."),
        ("Backend", "FastAPI, SQLAlchemy, MySQL, PyMySQL, python-jose, bcrypt."),
        ("AI and retrieval", "Groq, ChromaDB, all-MiniLM-L6-v2, prompts JSON."),
        ("Speech", "Browser MediaRecorder/AudioContext and faster-whisper."),
        ("Modes", "Normal: resume/adaptive. Syllabus: temporary material/topic-sequential."),
        ("Core API", "auth, resumes, interviews, users, speech under /api."),
        ("Core data", "users, resumes, interview_sessions, interview_questions, answers, answer_evaluations."),
    ],
    [1.85, 5.3],
)
h2("Final architecture summary")
code_block("User\n  -> React UI and AuthContext\n  -> Axios bearer-token requests\n  -> FastAPI routers and dependencies\n  -> service orchestration\n  -> MySQL persistence plus ChromaDB retrieval\n  -> SBERT embeddings and Groq generation/evaluation\n  -> adaptive or syllabus progression\n  -> persisted results, history, and Normal Mode analytics")
p("SmartInterview is a two-mode interview system sharing one persistence and API layer. Normal Mode is resume-driven and score-adaptive through deterministic Bloom logic. Syllabus Mode is material-driven and topic-sequential through temporary vector collections. Both modes persist interview artifacts and use the same answer-evaluation pipeline, while the current implementation has the noted Syllabus evaluation and UI boundaries.")

doc.core_properties.title = "SmartInterview Implementation Documentation"
doc.core_properties.subject = "Current SmartInterview implementation architecture and technical reference"
doc.core_properties.author = "SmartInterview Documentation"
doc.save(OUT)
print(OUT)
